"""Run Steady-State Waveform 2 Port headless against simulated equipment and check Excel output."""
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
os.environ['PI_ATE_AUTO_CONFIRM'] = '1'
time.sleep = lambda _s: None  # must run before test modules bind `from time import sleep`

from PySide2.QtWidgets import QApplication  # noqa: E402

app = QApplication.instance() or QApplication(sys.argv)

import openpyxl  # noqa: E402
from psu_tests.definitions import (TestConditions, LineRange, LoadRange, SoakTime,  # noqa: E402
                                   GeneralOptions, USBPDOptions, I2CTestParameters, TestStatus)
from psu_tests.test_type import TestTypes  # noqa: E402
from psu_tests.test_efficiency_2port_waveform import EfficiencyWaveformTest_2Port  # noqa: E402
from equipment.simulated_handler import SimulatedEquipmentHandler  # noqa: E402

MODES = {
    'A': "Option A: Proportional Sync (Both Ports 100% -> 10%)",
    'B': "Option B: Cross-Reg Matrix (Port 2 at Specified % Loads)",
}


class FakeParent:
    def __init__(self):
        self.run_settings = {'debug': False}
        self.equipment = SimulatedEquipmentHandler(self)


class FakeItem:
    def __init__(self, parent, conditions):
        self.parent = parent
        self.test_type_index = TestTypes.index(EfficiencyWaveformTest_2Port)
        self.test_conditions = conditions


def run_case(mode_key: str, out_dir: str) -> list:
    vin = [[115, 60], [230, 50]]
    loads = [100, 50, 10]
    conditions = TestConditions(
        name=EfficiencyWaveformTest_2Port.title,
        nominal_output_voltage_V=5, nominal_load_current_A=5, max_load_current_A=5,
        line_range=LineRange(name='TEST', vin_freq=vin, custom=True),
        load_range=LoadRange(name='TEST', load_range_pct=loads, custom=True),
        soak_time=SoakTime(name='VERIFY_SOAK', initial=2, line=3, load=4, integration=1),
        general_options=GeneralOptions(coupling='AC'),
        usbpd_options=USBPDOptions(),
        i2c_test_parameters=I2CTestParameters(
            params=[12.0, 2.0, 0.0, "100, 50", 3.0, 0, 0, 25.0, 0, 2],
            cbx_params=[MODES[mode_key], 'Yes', 'None (Fully Automated)', 'No']),
        unit_id='U7', ambient_temp=25.0,
        scope_channels={
            1: {
                'enabled': True,
                'name': 'Primary Vds',
                'derating': 650,
                'measurements': ['MAXimum', 'MINimum'],
            },
            2: {'enabled': True, 'name': 'Ids', 'measurements': ['MAXimum']},
            3: {'enabled': False, 'name': ''},
            4: {'enabled': False, 'name': ''},
        })

    statuses = []
    test = EfficiencyWaveformTest_2Port(FakeItem(FakeParent(), conditions))
    test.status_update.connect(statuses.append)
    test.output_folder_path = os.path.join(out_dir, f"mode_{mode_key}")
    test.run()

    errors = []
    if TestStatus.COMPLETE not in statuses:
        return [f"mode {mode_key}: test did not complete, statuses={statuses}"]

    expected_basename = (
        f"SSW 115-230VAC 5V 25W+24W Primary Vds+Ids")
    if test.data_filename != expected_basename:
        errors.append(
            f"mode {mode_key}: data_filename {test.data_filename!r} != {expected_basename!r}")
    if not os.path.isfile(test.data_file_path):
        errors.append(f"mode {mode_key}: missing workbook {test.data_file_path!r}")
    pngs = [
        rec.get('path') for rec in getattr(test, 'captured_records', [])
        if rec.get('path')]
    if pngs and not all('_P1_' in p and '_P2_' in p for p in pngs):
        errors.append(f"mode {mode_key}: PNG paths missing P1/P2 power suffix: {pngs[:2]!r}")

    n_points = len(test.load_pairs)
    wb = openpyxl.load_workbook(test.data_file_path)
    ws = wb[test.sheet_name]
    scope_section_col = test.scope_excel_col()
    hdr_row = test.excel_section_header_row
    col_row = test.excel_first_data_row
    if ws.cell(row=hdr_row, column=scope_section_col).value != 'Oscilloscope Measurements':
        errors.append(
            f"mode {mode_key}: scope section title row {hdr_row} col {scope_section_col} = "
            f"{ws.cell(row=hdr_row, column=scope_section_col).value!r}")
    if ws.cell(row=col_row, column=19).value != 'Efficiency (%)':
        errors.append(f"mode {mode_key}: S{col_row} = {ws.cell(row=col_row, column=19).value!r}")
    if ws.cell(row=col_row, column=scope_section_col).value != 'Trig Ch':
        errors.append(
            f"mode {mode_key}: scope header col {scope_section_col} = "
            f"{ws.cell(row=col_row, column=scope_section_col).value!r}")
    if ws.cell(row=hdr_row, column=20).value == 'Average Efficiency Limits':
        errors.append(f"mode {mode_key}: unexpected Average Efficiency Limits section at column T")
    if ws['A2'].value != 'Unit ID' or ws['B2'].value != 'U7':
        errors.append(f"mode {mode_key}: metadata A2/B2 = {ws['A2'].value!r}/{ws['B2'].value!r}")
    if ws['A9'].value != 'Soak profile' or ws['B9'].value != 'VERIFY_SOAK':
        errors.append(f"mode {mode_key}: soak profile = {ws['A9'].value!r}/{ws['B9'].value!r}")
    if ws['B10'].value != 2 or ws['B11'].value != 3 or ws['B12'].value != 4 or ws['B13'].value != 1:
        errors.append(
            f"mode {mode_key}: soak times B10-B13 = "
            f"{ws['B10'].value!r}/{ws['B11'].value!r}/{ws['B12'].value!r}/{ws['B13'].value!r}")
    if not test.sheet_name.startswith('SS_WF_2P_'):
        errors.append(f"mode {mode_key}: sheet name {test.sheet_name!r}")
    expected_scope = [
        'Trig Ch', 'Trig Level (V)', 'Primary Vds Max', 'Primary Vds Derating (%)',
        'Primary Vds Min', 'Ids Max',
    ]
    got_scope = [ws.cell(row=col_row, column=scope_section_col + i).value for i in range(len(expected_scope))]
    if got_scope != expected_scope:
        errors.append(f"mode {mode_key}: scope headers {got_scope}")

    data_start = col_row + 1
    for vin_idx in range(len(vin)):
        for r in range(data_start, data_start + n_points):
            eff = ws.cell(row=r, column=19).value
            vds = ws.cell(row=r, column=scope_section_col + 2).value
            eff_fmt = ws.cell(row=r, column=19).number_format
            if not isinstance(eff, (int, float)) or not isinstance(vds, (int, float)):
                errors.append(f"mode {mode_key}: row {r} eff={eff!r} vds={vds!r}")
            if eff_fmt != '#,##0.00':
                errors.append(f"mode {mode_key}: row {r} efficiency format {eff_fmt!r}")
            if ws.cell(row=r, column=scope_section_col).value != 'CH2':
                errors.append(
                    f"mode {mode_key}: row {r} trigger channel "
                    f"{ws.cell(row=r, column=scope_section_col).value!r}")
            last_scope_col = scope_section_col + len(expected_scope) - 1
            if ws.cell(row=r, column=last_scope_col + 1).value == 'Waveform File':
                errors.append(f"mode {mode_key}: Waveform File column should be omitted when embed=Yes")
        if vin_idx > 0:
            r0 = data_start
            if ws.cell(row=r0, column=1).value == 'Ch1 Load':
                errors.append(f"mode {mode_key}: repeated column header at row {r0} for Vin block {vin_idx}")
            if ws.cell(row=r0, column=19).value == 'Efficiency (%)':
                errors.append(f"mode {mode_key}: repeated efficiency header at row {r0}")
        data_start += n_points

    wf_sheet = f"{test.short_title}_Waveforms"
    if wf_sheet not in wb.sheetnames:
        errors.append(f"mode {mode_key}: no {wf_sheet} sheet")
    elif len(wb[wf_sheet]._images) != n_points * len(vin):
        errors.append(f"mode {mode_key}: {len(wb[wf_sheet]._images)} embedded images, expected {n_points * len(vin)}")
    else:
        ws_wf = wb[wf_sheet]
        if ws_wf['B4'].value and 'Primary Vds' in str(ws_wf['B4'].value):
            errors.append(f"mode {mode_key}: waveform caption should not include scope measurements")
        found_param_hdr = any(
            ws_wf.cell(row=row, column=2).value == 'Parameter'
            for row in range(4, min(ws_wf.max_row, 80)))
        if not found_param_hdr:
            errors.append(f"mode {mode_key}: waveforms sheet missing detail Parameter/Value table")
        if ws_wf._images and getattr(ws_wf._images[0], 'width', 0) <= 420:
            errors.append(f"mode {mode_key}: embedded image width {ws_wf._images[0].width} (expect >420 native scale)")

    table = test.test_data_table
    if len(table.header) != 19 + len(expected_scope):
        errors.append(
            f"mode {mode_key}: results table has {len(table.header)} columns, "
            f"expected {19 + len(expected_scope)}")
    data_rows = [r for r in table.data if any(v != '' for v in r)]
    if len(data_rows) != n_points * len(vin):
        errors.append(f"mode {mode_key}: results table has {len(data_rows)} data rows")

    print(f"mode {mode_key}: {n_points} load pairs x {len(vin)} line steps -> {test.data_file_path}")
    return errors


def main() -> int:
    out_dir = tempfile.mkdtemp(prefix="eff2p_wf_")
    errors = run_case('A', out_dir) + run_case('B', out_dir)
    if errors:
        print("FAIL")
        for e in errors:
            print(" ", e)
        return 1
    print("OK Steady-State Waveform 2 Port simulated runs passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
