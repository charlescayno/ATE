"""Run Startup Waveform 2 Port headless against simulated equipment and check Excel output."""
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.chdir(ROOT)
os.environ['PI_ATE_AUTO_CONFIRM'] = '1'
time.sleep = lambda _s: None  # must run before test modules bind `from time import sleep`

from PySide2.QtWidgets import QApplication  # noqa: E402

app = QApplication.instance() or QApplication(sys.argv)

import openpyxl  # noqa: E402
from psu_tests.definitions import (TestConditions, LineRange, LoadRange, SoakTime,  # noqa: E402
                                   GeneralOptions, USBPDOptions, I2CTestParameters, TestStatus)
from psu_tests.test_type import TestTypes  # noqa: E402
from psu_tests.test_startup_2port_waveform import StartupWaveformTest_2Port  # noqa: E402
from equipment.simulated_handler import SimulatedEquipmentHandler  # noqa: E402

MODES = {
    'A': "Option A: Proportional Sync (Both Ports 100% -> 10%)",
}


class FakeParent:
    def __init__(self):
        self.run_settings = {'debug': False}
        self.equipment = SimulatedEquipmentHandler(self)


class FakeItem:
    def __init__(self, parent, conditions):
        self.parent = parent
        self.test_type_index = TestTypes.index(StartupWaveformTest_2Port)
        self.test_conditions = conditions


def run_case(mode_key: str, out_dir: str) -> list:
    vin = [[115, 60], [230, 50]]
    loads = [100, 50]
    conditions = TestConditions(
        name=StartupWaveformTest_2Port.title,
        nominal_output_voltage_V=5, nominal_load_current_A=5, max_load_current_A=5,
        line_range=LineRange(name='TEST', vin_freq=vin, custom=True),
        load_range=LoadRange(name='TEST', load_range_pct=loads, custom=True),
        soak_time=SoakTime(name='VERIFY_SOAK', initial=2, line=3, load=4, integration=1),
        general_options=GeneralOptions(coupling='AC'),
        usbpd_options=USBPDOptions(),
        i2c_test_parameters=I2CTestParameters(
            params=[12.0, 2.0, 0.0, "100, 50", 3.0, 0, 0, 25.0, 0, 2],
            cbx_params=[MODES[mode_key], 'Yes', 'None (Fully Automated)', 'No']),
        unit_id='U7', ambient_temp=25.0,
        scope_channels={
            1: {
                'enabled': True,
                'name': 'Primary Vds',
                'derating': 650,
                'measurements': ['MAXimum', 'MINimum'],
            },
            2: {'enabled': True, 'name': 'Ids', 'measurements': ['MAXimum']},
            3: {'enabled': False, 'name': ''},
            4: {'enabled': False, 'name': ''},
        })

    statuses = []
    test = StartupWaveformTest_2Port(FakeItem(FakeParent(), conditions))
    test.status_update.connect(statuses.append)
    test.output_folder_path = os.path.join(out_dir, f"mode_{mode_key}")
    test.run()

    errors = []
    if TestStatus.COMPLETE not in statuses:
        return [f"mode {mode_key}: test did not complete, statuses={statuses}"]

    expected_basename = (
        f"SUW 115-230VAC 5V 25W+24W Primary Vds+Ids")
    if test.data_filename != expected_basename:
        errors.append(
            f"mode {mode_key}: data_filename {test.data_filename!r} != {expected_basename!r}")
    if not os.path.isfile(test.data_file_path):
        errors.append(f"mode {mode_key}: missing workbook {test.data_file_path!r}")
    if not test.sheet_name.startswith('SU_WF_2P_'):
        errors.append(f"mode {mode_key}: sheet name {test.sheet_name!r}")
    pngs = [
        rec.get('path') for rec in getattr(test, 'captured_records', [])
        if rec.get('path')]
    if pngs and not all('_P1_' in p and '_P2_' in p and p.endswith('_SUW.png') for p in pngs):
        errors.append(f"mode {mode_key}: PNG paths missing P1/P2/_SUW suffix: {pngs[:2]!r}")
    if test.auto_find_trigger:
        errors.append(f"mode {mode_key}: auto_find_trigger should be False")

    n_points = len(test.load_pairs)
    wb = openpyxl.load_workbook(test.data_file_path)
    ws = wb[test.sheet_name]
    scope_section_col = test.scope_excel_col()
    col_row = test.excel_first_data_row
    data_start = col_row + 1
    data_rows = n_points * len(vin)
    if ws.cell(row=data_start, column=19).value is None:
        errors.append(f"mode {mode_key}: missing efficiency at row {data_start}")

    wf_sheet = f"{test.short_title}_Waveforms"
    if wf_sheet not in wb.sheetnames:
        errors.append(f"mode {mode_key}: no {wf_sheet} sheet")
    elif len(wb[wf_sheet]._images) != data_rows:
        errors.append(
            f"mode {mode_key}: {len(wb[wf_sheet]._images)} embedded images, expected {data_rows}")

    print(f"mode {mode_key}: {n_points} load pairs x {len(vin)} line steps -> {test.data_file_path}")
    return errors


def main() -> int:
    out_dir = tempfile.mkdtemp(prefix="su_wf_2p_")
    errors = run_case('A', out_dir)
    if errors:
        print("FAIL")
        for e in errors:
            print(" ", e)
        return 1
    print("OK Startup Waveform 2 Port simulated runs passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
