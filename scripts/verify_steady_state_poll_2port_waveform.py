"""Headless run of Steady-State Poll Waveform 2 Port against simulated equipment."""
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path[:] = [str(ROOT)] + [
    p for p in sys.path[1:]
    if p and "PI_ATE_Marketing_Release" not in p.replace("/", os.sep)
]
os.chdir(ROOT)
os.environ['PI_ATE_AUTO_CONFIRM'] = '1'
time.sleep = lambda _s: None

from PySide2.QtWidgets import QApplication  # noqa: E402

app = QApplication.instance() or QApplication(sys.argv)

import datetime

import openpyxl  # noqa: E402
from psu_tests.definitions import (TestConditions, LineRange, LoadRange, SoakTime,  # noqa: E402
                                   GeneralOptions, USBPDOptions, I2CTestParameters, TestStatus,
                                   test_control_flags)
from psu_tests.test_type import TestTypes  # noqa: E402
from psu_tests.test_steady_state_poll_2port_waveform import SteadyStatePollWaveformTest_2Port  # noqa: E402
from equipment.simulated_handler import SimulatedEquipmentHandler  # noqa: E402
from page_controls.eff2p_waveform_advanced import Eff2PWaveformAdvancedDialog  # noqa: E402
from page_controls.test_sequence_flow_text import build_flow_text_for_title  # noqa: E402
from page_controls.poll_sequence_diagram import render_poll_sequence_diagram_png  # noqa: E402
from PySide2.QtWidgets import QLineEdit  # noqa: E402

OPTION_C = "Option C: Fixed Aux / Swept Main (Port 2 Fixed Load)"
POLL = SteadyStatePollWaveformTest_2Port


def check(errors, name, ok, detail=""):
    if ok:
        print(f"OK  {name}")
    else:
        print(f"FAIL {name} {detail}")
        errors.append(f"{name}: {detail}")


def test_effective_integration(errors):
    check(errors, "eff int from soak only", POLL.effective_integration_s(10, 300) == 10.0)
    check(errors, "eff int soak larger", POLL.effective_integration_s(300, 10) == 300.0)
    check(errors, "eff int ignores poll wait", POLL.effective_integration_s(5, 300) == 5.0)
    check(errors, "eff int zero soak", POLL.effective_integration_s(0, 300) == 0.0)


def test_fan_cooling(errors):
    check(errors, "fan param default off", not POLL.fan_cooling_enabled_from_params(
        list(POLL.tc_default.i2c_test_parameters.param)))
    check(errors, "fan param on", POLL.fan_cooling_enabled_from_params(
        [0.0] * POLL.FAN_COOLING_PARAM_INDEX + [1.0]))
    soak = SoakTime(name="TST", initial=0, line=0, load=0, integration=0, custom=False)
    params = [12.0, 2.0, 0.0, 0, 1.0, 0, 0, 25.0, 0, 1, 1.0, 3.0, 1.0]
    conditions = TestConditions(
        name=POLL.title,
        nominal_output_voltage_V=5,
        nominal_load_current_A=5,
        max_load_current_A=5,
        line_range=LineRange(name="115", vin_freq=[[115, 60]], custom=True),
        load_range=LoadRange(name="100", load_range_pct=[100], custom=True),
        soak_time=soak,
        general_options=GeneralOptions(),
        usbpd_options=USBPDOptions(),
        i2c_test_parameters=I2CTestParameters(
            params=params,
            cbx_params=[OPTION_C, 'Yes', 'None (Fully Automated)', 'No']),
        unit_id='P1', ambient_temp=25.0,
        scope_channels={
            1: {'enabled': True, 'name': 'Primary Vds', 'derating': 650, 'measurements': ['MAXimum']},
            2: {'enabled': True, 'name': 'Ids', 'measurements': ['MAXimum']},
            3: {'enabled': False, 'name': ''},
            4: {'enabled': False, 'name': ''},
        })
    parent = FakeParent()
    hmp = parent.equipment.dc_source_aux
    hmp._hmp_channels[2]['on'] = 0
    test = POLL(FakeItem(parent, conditions))
    test.setup_equipment()
    ch2 = hmp._hmp_channels[2]
    check(errors, "fan ch2 on at start", ch2['on'] == 1, ch2)
    check(errors, "fan ch2 voltage", ch2['v'] == POLL.FAN_COOLING_VOLTAGE_V, ch2['v'])
    check(errors, "fan ch2 current", ch2['i'] == POLL.FAN_COOLING_CURRENT_A, ch2['i'])
    test._shutdown_hmp_fan_cooling()
    check(errors, "fan off after teardown", ch2['on'] == 0, ch2)


def test_sample_math(errors):
    check(errors, "legacy duration migration", POLL.read_samples_per_point(
        [0.0] * 10 + [30.0, 1200.0], 30.0) == 41)
    check(errors, "direct sample count", POLL.read_samples_per_point(
        [0.0] * 10 + [30.0, 5.0], 30.0) == 5)
    check(errors, "samples min 1", POLL.read_samples_per_point(
        [0.0] * 10 + [30.0, 0.0], 30.0) == 1)
    check(errors, "legacy sample_count_for", POLL.sample_count_for(5, 900) == 181)
    params = list(POLL.tc_default.i2c_test_parameters.param)
    check(
        errors, "default poll params",
        len(params) > POLL.POLL_SAMPLES_PARAM_INDEX
        and float(params[POLL.POLL_INTERVAL_PARAM_INDEX]) == 30.0
        and float(params[POLL.POLL_SAMPLES_PARAM_INDEX]) == 3.0,
        params)
    check(errors, "registered in TestTypes", POLL in TestTypes)
    check(errors, "does not shift waveform index 0",
          TestTypes[0].title == "Steady-State Waveform 2 Port")


def test_dialog(errors):
    interval = QLineEdit()
    samples = QLineEdit()
    interval.setText("30")
    samples.setText("5")
    hidden = [QLineEdit() for _ in range(10)]
    dlg = Eff2PWaveformAdvancedDialog(
        None, hidden, scope_channels={1: {'enabled': True, 'name': 'Primary Vds'}},
        for_poll=True, poll_interval_le=interval, poll_duration_le=samples)
    check(errors, "dialog poll fields",
          dlg.le_poll_interval is not None and dlg.le_poll_duration.text() == "5")
    dlg.le_poll_interval.setText("0.5")
    dlg.le_poll_duration.setText("4")
    dlg.apply_to_lineedits()
    check(errors, "dialog writes poll storage", interval.text() == "0.5" and samples.text() == "4",
          f"{interval.text()} {samples.text()}")
    flow = build_flow_text_for_title(
        POLL.title, poll_interval_s=0.5, sample_count=4)
    check(errors, "flow text", flow and "0.5 s" in flow and "SS_POLL_2P" in flow, flow[:80] if flow else None)
    png = render_poll_sequence_diagram_png(
        n_samples=4, poll_wait_s=0.5, integration_s=10.0)
    check(errors, "sequence diagram png", len(png) > 8000, len(png))


class FakeParent:
    def __init__(self):
        self.run_settings = {'debug': False}
        self.equipment = SimulatedEquipmentHandler(self)


class FakeItem:
    def __init__(self, parent, conditions):
        self.parent = parent
        self.test_conditions = conditions
        self.test_type_index = TestTypes.index(POLL)


def test_run_complete(errors):
    soak = SoakTime(name="TST", initial=0, line=0, load=0, integration=0, custom=False)
    conditions = TestConditions(
        name=POLL.title,
        nominal_output_voltage_V=5,
        nominal_load_current_A=5,
        max_load_current_A=5,
        line_range=LineRange(name="115", vin_freq=[[115, 60]], custom=True),
        load_range=LoadRange(name="100", load_range_pct=[100], custom=True),
        soak_time=soak,
        general_options=GeneralOptions(),
        usbpd_options=USBPDOptions(),
        i2c_test_parameters=I2CTestParameters(
            params=[12.0, 2.0, 0.0, 0, 1.0, 0, 0, 25.0, 0, 1, 1.0, 3.0],
            cbx_params=[OPTION_C, 'Yes', 'None (Fully Automated)', 'No']),
        unit_id='P1', ambient_temp=25.0,
        scope_channels={
            1: {'enabled': True, 'name': 'Primary Vds', 'derating': 650, 'measurements': ['MAXimum']},
            2: {'enabled': True, 'name': 'Ids', 'measurements': ['MAXimum']},
            3: {'enabled': False, 'name': ''},
            4: {'enabled': False, 'name': ''},
        })
    statuses = []
    with tempfile.TemporaryDirectory() as out_dir:
        test = POLL(FakeItem(FakeParent(), conditions))
        test.status_update.connect(statuses.append)
        test.output_folder_path = out_dir
        check(errors, "queue estimate includes poll",
              test.total_time >= 2.0, test.total_time)
        test.run()
        check(errors, "run complete", TestStatus.COMPLETE in statuses, statuses)
        expected_name = "SSP 115VAC 5V 25W+24W Primary Vds+Ids"
        check(errors, "workbook name", test.data_filename == expected_name, test.data_filename)
        check(errors, "workbook exists", os.path.isfile(test.data_file_path), test.data_file_path)
        pngs = [rec.get('path') for rec in test.captured_records if rec.get('path')]
        check(errors, "screenshot per sample",
              len(pngs) == 3 and all(p.endswith('_SSP_S{}.png'.format(i)) for i, p in enumerate(pngs, 1)),
              pngs)
        check(errors, "capture record per sample",
              len(test.captured_records) == 3
              and all(rec.get('main_sheet_row') for rec in test.captured_records),
              test.captured_records)

        wb = openpyxl.load_workbook(test.data_file_path)
        check(errors, "sheet name", test.sheet_name in wb.sheetnames and test.sheet_name.startswith('SS_POLL_2P_'),
              test.sheet_name)
        ws = wb[test.sheet_name]
        check(errors, "poll metadata",
              ws['D2'].value == 'Poll wait between samples (s)' and ws['E2'].value == 1
              and ws['E3'].value == 3 and ws['D3'].value == 'Samples per load point',
              (ws['D2'].value, ws['E2'].value, ws['D3'].value, ws['E3'].value))
        check(errors, "section title", ws['T14'].value == 'Timed Poll', ws['T14'].value)
        hdr = test.excel_first_data_row
        poll_headers = [ws.cell(row=hdr, column=c).value for c in range(20, 26)]
        check(errors, "poll headers",
              poll_headers == list(POLL.POLL_COLUMNS), poll_headers)
        scope_col = test.scope_excel_col()
        check(errors, "scope starts after poll columns", scope_col == 26 and ws.cell(row=hdr, column=scope_col).value == 'Trig Ch',
              (scope_col, ws.cell(row=hdr, column=scope_col).value))
        for i, (sample, elapsed) in enumerate(((1, 0), (2, 1), (3, 2)), start=1):
            row = hdr + i
            got_s = ws.cell(row=row, column=20).value
            got_e = ws.cell(row=row, column=21).value
            t_start = ws.cell(row=row, column=22).value
            t_end = ws.cell(row=row, column=23).value
            t_sampled = ws.cell(row=row, column=24).value
            eff = ws.cell(row=row, column=19).value
            vds = ws.cell(row=row, column=scope_col + 2).value
            times_ok = (
                isinstance(t_start, datetime.datetime)
                and isinstance(t_end, datetime.datetime)
                and isinstance(t_sampled, datetime.datetime)
                and t_end >= t_start
                and t_sampled >= t_start
            )
            check(errors, f"row {row} sample",
                  got_s == sample and float(got_e) == float(elapsed)
                  and times_ok
                  and isinstance(eff, (int, float)) and isinstance(vds, (int, float)),
                  (got_s, got_e, t_start, t_end, t_sampled, eff, vds))
        row1_start = ws.cell(row=hdr + 1, column=22).value
        row1_end = ws.cell(row=hdr + 1, column=23).value
        row3_start = ws.cell(row=hdr + 3, column=22).value
        row3_end = ws.cell(row=hdr + 3, column=23).value
        check(errors, "time window same per load point",
              row1_start == row3_start and row1_end == row3_end,
              (row1_start, row3_start, row1_end, row3_end))
        window_s = (row1_end - row1_start).total_seconds()
        check(errors, "time end after time start",
              window_s > 0, window_s)
        sampled = [ws.cell(row=hdr + i, column=24).value for i in (1, 2, 3)]
        check(errors, "time sampled monotonic",
              sampled[0] <= sampled[1] <= sampled[2], sampled)
        fmt = ws.cell(row=hdr + 1, column=22).number_format
        check(errors, "time start excel format",
              fmt == POLL.POLL_DATETIME_EXCEL_FMT, fmt)
        iv1 = ws.cell(row=hdr + 1, column=25).value
        iv2 = ws.cell(row=hdr + 2, column=25).value
        check(errors, "interval sample1 zero", float(iv1) == 0.0, iv1)
        check(errors, "interval sample2 non-negative", float(iv2) >= 0.0, iv2)
        if sampled[1] and sampled[0]:
            expected_iv2 = round((sampled[1] - sampled[0]).total_seconds(), 2)
            check(errors, "interval matches time sampled delta",
                  abs(float(iv2) - expected_iv2) < 0.01, (iv2, expected_iv2))
        wf = f"{test.short_title}_Waveforms"
        check(errors, "embedded image per sample",
              wf in wb.sheetnames and len(wb[wf]._images) == 3,
              wf in wb.sheetnames and len(wb[wf]._images))
        check(errors, "results table rows",
              len([r for r in test.test_data_table.data if any(v != '' for v in r)]) == 3)


def test_abort_saves_excel(errors):
    """Stop after two samples; those rows must already be in the workbook."""
    test_control_flags['StopTest'] = False
    test_control_flags['SkipTest'] = False
    soak = SoakTime(name="TST", initial=0, line=0, load=0, integration=0, custom=False)
    conditions = TestConditions(
        name=POLL.title,
        nominal_output_voltage_V=5,
        nominal_load_current_A=5,
        max_load_current_A=5,
        line_range=LineRange(name="115", vin_freq=[[115, 60]], custom=True),
        load_range=LoadRange(name="100", load_range_pct=[100], custom=True),
        soak_time=soak,
        general_options=GeneralOptions(),
        usbpd_options=USBPDOptions(),
        i2c_test_parameters=I2CTestParameters(
            params=[12.0, 2.0, 0.0, 0, 1.0, 0, 0, 25.0, 0, 1, 5.0, 5.0],
            cbx_params=[OPTION_C, 'Yes', 'None (Fully Automated)', 'No']),
        unit_id='P1', ambient_temp=25.0,
        scope_channels={
            1: {'enabled': True, 'name': 'Primary Vds', 'derating': 650, 'measurements': ['MAXimum']},
            2: {'enabled': False, 'name': ''},
            3: {'enabled': False, 'name': ''},
            4: {'enabled': False, 'name': ''},
        })
    statuses = []
    try:
        with tempfile.TemporaryDirectory() as out_dir:
            test = POLL(FakeItem(FakeParent(), conditions))
            test.status_update.connect(statuses.append)
            test.output_folder_path = out_dir
            original = test._append_measurement_row

            def stop_after_two(scope_row):
                original(scope_row)
                if test.poll_sample_index >= 2:
                    test_control_flags['StopTest'] = True

            test._append_measurement_row = stop_after_two
            test.run()
            check(errors, "abort status", TestStatus.STOPPED in statuses, statuses)
            check(errors, "abort workbook exists", os.path.isfile(test.data_file_path))
            wb = openpyxl.load_workbook(test.data_file_path)
            ws = wb[test.sheet_name]
            hdr = test.excel_first_data_row
            samples = [ws.cell(row=hdr + i, column=20).value for i in (1, 2, 3)]
            elapsed = [ws.cell(row=hdr + i, column=21).value for i in (1, 2, 3)]
            check(errors, "abort kept two samples",
                  samples[:2] == [1, 2] and [float(x) for x in elapsed[:2]] == [0.0, 5.0]
                  and samples[2] in (None, ""),
                  (samples, elapsed))
            t0 = ws.cell(row=hdr + 1, column=24).value
            t1 = ws.cell(row=hdr + 2, column=24).value
            abort_end = ws.cell(row=hdr + 2, column=23).value
            abort_start = ws.cell(row=hdr + 2, column=22).value
            check(errors, "abort rows have timestamps",
                  isinstance(t0, datetime.datetime) and isinstance(t1, datetime.datetime)
                  and isinstance(abort_start, datetime.datetime)
                  and isinstance(abort_end, datetime.datetime),
                  (t0, t1, abort_start, abort_end))
            check(errors, "abort capture per sample",
                  len(test.captured_records) == 2
                  and test.captured_records[0]['path'].endswith('_SSP_S1.png')
                  and test.captured_records[1]['path'].endswith('_SSP_S2.png'),
                  test.captured_records)
    finally:
        test_control_flags['StopTest'] = False
        test_control_flags['SkipTest'] = False


def main():
    errors = []
    test_effective_integration(errors)
    test_sample_math(errors)
    test_fan_cooling(errors)
    test_dialog(errors)
    test_run_complete(errors)
    test_abort_saves_excel(errors)
    if errors:
        print("FAIL")
        for e in errors:
            print(f"  {e}")
        sys.exit(1)
    print(f"OK {POLL.title}")


if __name__ == "__main__":
    main()
