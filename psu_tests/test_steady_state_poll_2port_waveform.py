"""Steady-state 2-port waveform test that polls meters for a timed window."""
import datetime
import math
import time
from copy import deepcopy
from time import sleep

import openpyxl
from openpyxl.styles import Font as CellFont

from psu_tests.definitions import TestConditions, I2CTestParameters, test_control_flags, TestFailed
from psu_tests.test_object_imports import TestStopped, TestSkipped
from psu_tests.test_efficiency_2port_waveform import EfficiencyWaveformTest_2Port


class SteadyStatePollWaveformTest_2Port(EfficiencyWaveformTest_2Port):
    """
    Steady-State Poll Waveform 2 Port.

    Same line/load sweep and one scope screenshot per point as Steady-State
    Waveform 2 Port. Each sample: integrate Pin (param 10) → meters → scope capture (PNG)
    → Excel row → repeat for N samples. Next sample integration starts after previous scope finishes.
    """

    title = "Steady-State Poll Waveform 2 Port"
    short_title = "SS_POLL_2P"
    PNG_SUFFIX = "_SSP"
    POLL_COLUMNS = (
        "Sample",
        "Elapsed (s)",
        "Time Start",
        "Time End",
        "Time Sampled",
        "Time Integration Start",
        "Time Integration End",
        "Interval (s)",
    )
    POLL_DATETIME_EXCEL_FMT = "yyyy-mm-dd hh:mm:ss"
    POLL_INTERVAL_PARAM_INDEX = 10
    POLL_SAMPLES_PARAM_INDEX = 11
    POLL_DURATION_PARAM_INDEX = 11  # legacy alias (stored value is sample count)
    FAN_COOLING_PARAM_INDEX = 12
    FAN_COOLING_VOLTAGE_V = 12.0
    FAN_COOLING_CURRENT_A = 4.0
    HMP_FAN_CHANNEL = 2
    POLL_INTERVAL_DEFAULT_S = 30.0
    POLL_READ_DURATION_GUESS_S = 0.75
    POLL_SAMPLES_DEFAULT = 3
    POLL_DURATION_DEFAULT_S = 20 * 60.0  # legacy duration for migration only
    LEGACY_POLL_DURATION_MIN_S = 60.0
    POLL_SCOPE_CAPTURE_GUESS_S = 12.0
    POLL_DEFAULT_SCOPE_CHANNELS = {
        1: {
            'enabled': True,
            'name': 'VDS_PRI',
            'measurements': ['MAXimum'],
            'derating': 1700.0,
            'cursor': {'enabled': False, 'type': 'VERTical'},
        },
        2: {'enabled': False, 'name': '', 'measurements': [], 'derating': None, 'cursor': {}},
        3: {'enabled': False, 'name': '', 'measurements': [], 'derating': None, 'cursor': {}},
        4: {'enabled': False, 'name': '', 'measurements': [], 'derating': None, 'cursor': {}},
    }

    i2c_ui_definitions = EfficiencyWaveformTest_2Port.i2c_ui_definitions
    tc_default = None

    @classmethod
    def _build_tc_default(cls):
        tc = EfficiencyWaveformTest_2Port.tc_default
        i2c = tc.i2c_test_parameters
        params = list(getattr(i2c, 'param', []) or [])
        while len(params) <= cls.POLL_SAMPLES_PARAM_INDEX:
            params.append(0.0)
        params[cls.POLL_INTERVAL_PARAM_INDEX] = cls.POLL_INTERVAL_DEFAULT_S
        params[cls.POLL_SAMPLES_PARAM_INDEX] = float(cls.POLL_SAMPLES_DEFAULT)
        while len(params) <= cls.FAN_COOLING_PARAM_INDEX:
            params.append(0.0)
        params[cls.FAN_COOLING_PARAM_INDEX] = 0.0
        cbx = list(getattr(i2c, 'cbx_param', []) or [])
        while len(cbx) < 4:
            cbx.append('')
        if not str(cbx[0]).strip():
            cbx[0] = "Option C: Fixed Aux / Swept Main (Port 2 Fixed Load)"
        cbx[2] = "None (Fully Automated)"
        cbx[3] = "No"
        return TestConditions(
            name=cls.title,
            nominal_output_voltage_V=tc.nominal_output_voltage_V,
            nominal_load_current_A=tc.nominal_load_current_A,
            max_load_current_A=tc.max_load_current_A,
            line_range=tc.line_range,
            load_range=tc.load_range,
            soak_time=tc.soak_time,
            general_options=tc.general_options,
            usbpd_options=tc.usbpd_options,
            line_ramp_settings=tc.line_ramp_settings,
            i2c_test_parameters=I2CTestParameters(
                params=params,
                cbx_params=cbx,
            ),
            unit_id=tc.unit_id,
            ambient_temp=tc.ambient_temp,
            scope_channels=deepcopy(cls.POLL_DEFAULT_SCOPE_CHANNELS),
        )

    @staticmethod
    def sample_count_for(interval_s, duration_s) -> int:
        """Samples at t=0, t=interval, ... while t <= duration. At least one."""
        try:
            interval = float(interval_s)
            duration = float(duration_s)
        except (TypeError, ValueError):
            return 1
        if interval <= 0 or duration <= 0:
            return 1
        return int(math.floor((duration + 1e-9) / interval)) + 1

    def _sample_count(self) -> int:
        return int(getattr(self, 'poll_samples_n', self.POLL_SAMPLES_DEFAULT))

    @classmethod
    def read_samples_per_point(cls, params, poll_interval_s=None) -> int:
        """Param 11 = sample count; values >= LEGACY_POLL_DURATION_MIN_S are old duration (s)."""
        interval = poll_interval_s
        if interval is None:
            interval = cls._read_poll_param(
                params, cls.POLL_INTERVAL_PARAM_INDEX, cls.POLL_INTERVAL_DEFAULT_S)
        raw = cls._read_poll_param(
            params, cls.POLL_SAMPLES_PARAM_INDEX, float(cls.POLL_SAMPLES_DEFAULT))
        if raw >= cls.LEGACY_POLL_DURATION_MIN_S and interval > 0:
            return cls.sample_count_for(interval, raw)
        return max(1, int(round(raw)))

    @staticmethod
    def effective_integration_s(soak_integration_s, poll_interval_s=None) -> float:
        """Pin integration window = poll param 10 (soak integration ignored for poll Pin)."""
        del soak_integration_s
        try:
            t = float(poll_interval_s or 0)
        except (TypeError, ValueError):
            t = 0.0
        if t < 0:
            t = 0.0
        return t

    def _effective_integration_s(self) -> float:
        return self.effective_integration_s(
            0,
            getattr(self, 'poll_interval_s', self.POLL_INTERVAL_DEFAULT_S),
        )

    def _configure_poll_power_meter_integration(self):
        timer_s = self._effective_integration_s()
        for pm, use_poll_timer in (
            (self.power_meter_source, True),
            (self.power_meter_load, False),
            (self.power_meter_load_2, False),
        ):
            if pm is None:
                continue
            t = timer_s if use_poll_timer else 0
            if not use_poll_timer and getattr(pm, 'integration_disabled', lambda: False)():
                t = 0
            pm.integration_settings(mode="NORMAL", timer_s=t)

    @classmethod
    def fan_cooling_enabled_from_params(cls, params) -> bool:
        raw = cls._read_poll_param(params, cls.FAN_COOLING_PARAM_INDEX, 0.0)
        return raw >= 0.5

    def _hmp_for_fan_cooling(self):
        eq = self.equipment
        if hasattr(eq, 'prefer_hmp_as_aux'):
            eq.prefer_hmp_as_aux()
        aux = getattr(eq, 'dc_source_aux', None)
        if aux is not None and aux.is_hmp_series():
            return aux
        primary = getattr(eq, 'dc_source', None)
        input_supply = getattr(self, 'input_supply', None)
        for dc in getattr(eq, 'dc_sources', []) or []:
            if dc is None or not dc.is_hmp_series():
                continue
            if dc is input_supply or dc is primary:
                continue
            return dc
        return None

    def _apply_hmp_fan_cooling_preflight(self):
        """Before test loop: fan ON when enabled, else ensure HMP CH2 is OFF."""
        self._fan_cooling_active = False
        hmp = self._hmp_for_fan_cooling()
        ch = self.HMP_FAN_CHANNEL
        if self.enable_fan_cooling:
            if hmp is None:
                raise TestFailed(
                    "Enable Fan Cooling is checked but no HMP aux DC is assigned "
                    "(Equipment Setup → Aux DC → HMP2030)."
                )
            v, a = self.FAN_COOLING_VOLTAGE_V, self.FAN_COOLING_CURRENT_A
            self.update_status_log(
                f"HMP CH{ch} fan cooling: {v:g} V, {a:g} A limit, output ON",
                upcoming="Prepare test",
            )
            hmp.channel_set_voltage(ch, v)
            hmp.channel_set_current_limit(ch, a)
            hmp.channel_turn_on(ch)
            self._fan_cooling_active = True
            print(
                f"[POLL] Fan cooling ON — {getattr(hmp, 'description', hmp)} "
                f"CH{ch} {v:g} V / {a:g} A"
            )
        elif hmp is not None:
            hmp.channel_turn_off(ch)
            print(f"[POLL] Fan cooling disabled — HMP CH{ch} output OFF")

    def _shutdown_hmp_fan_cooling(self):
        hmp = self._hmp_for_fan_cooling()
        if hmp is None:
            return
        ch = self.HMP_FAN_CHANNEL
        hmp.channel_turn_off(ch)
        self._fan_cooling_active = False
        print(f"[POLL] Fan cooling OFF — HMP CH{ch} output OFF")

    def run(self):
        try:
            super().run()
        finally:
            if getattr(self, '_fan_cooling_active', False):
                self._shutdown_hmp_fan_cooling()

    def setup_equipment(self):
        super().setup_equipment()
        self._apply_hmp_fan_cooling_preflight()
        self._configure_poll_power_meter_integration()
        eff = self._effective_integration_s()
        if eff > 0:
            print(
                f"[POLL] Pin integration per sample: {eff:g} s "
                f"(I2C param {self.POLL_INTERVAL_PARAM_INDEX})"
            )

    @classmethod
    def _read_poll_param(cls, params, index, default) -> float:
        try:
            if params and len(params) > index and str(params[index]).strip() != '':
                val = float(params[index])
                if val < 0:
                    return default
                return val
        except (TypeError, ValueError):
            pass
        return default

    def unpack_test_item(self):
        super().unpack_test_item()
        i2c = getattr(self.test_conditions, 'i2c_test_parameters', None)
        params = list(getattr(i2c, 'param', []) or [])
        self.poll_interval_s = self._read_poll_param(
            params, self.POLL_INTERVAL_PARAM_INDEX, self.POLL_INTERVAL_DEFAULT_S)
        self.poll_samples_n = self.read_samples_per_point(params, self.poll_interval_s)
        self.poll_sample_index = 1
        self.poll_elapsed_s = 0.0
        self.poll_point_start_dt = None
        self.poll_point_end_dt = None
        self.poll_row_sampled_dt = None
        self.poll_integ_start_dt = None
        self.poll_integ_end_dt = None
        self.poll_prev_sampled_dt = None
        self._poll_repeat_read_duration_s = self.POLL_READ_DURATION_GUESS_S
        self.enable_fan_cooling = self.fan_cooling_enabled_from_params(params)
        self._fan_cooling_active = False

    def build_detailed_output_basename(self) -> str:
        return (
            f"SSP {self._line_vin_filename_part()} {self.vout_V:g}V "
            f"{self._dual_port_power_filename_part()} {self._scope_names_filename_part()}"
        )

    def define_data_header(self):
        return list(super().define_data_header()) + list(self.POLL_COLUMNS)

    def scope_excel_col(self) -> int:
        return super().scope_excel_col() + len(self.POLL_COLUMNS)

    def process_data_row(self):
        row = super().process_data_row()
        if (
            self.poll_row_sampled_dt is not None
            and self.poll_prev_sampled_dt is not None
        ):
            interval_s = round(
                (self.poll_row_sampled_dt - self.poll_prev_sampled_dt).total_seconds(),
                2,
            )
        else:
            interval_s = 0.0
        return row + [
            int(self.poll_sample_index),
            round(float(self.poll_elapsed_s), 2),
            self.poll_point_start_dt,
            self.poll_point_end_dt,
            self.poll_row_sampled_dt,
            self.poll_integ_start_dt,
            self.poll_integ_end_dt,
            interval_s,
        ]

    def _poll_first_excel_col(self) -> int:
        return self.scope_excel_col() - len(self.POLL_COLUMNS)

    @staticmethod
    def _excel_num_format_for_header(header: str) -> str:
        if header == 'Sample':
            return '0'
        if header in ('Elapsed (s)', 'Interval (s)'):
            return '#,##0.00'
        if header in (
            'Time Start',
            'Time End',
            'Time Sampled',
            'Time Integration Start',
            'Time Integration End',
        ):
            return SteadyStatePollWaveformTest_2Port.POLL_DATETIME_EXCEL_FMT
        return EfficiencyWaveformTest_2Port._excel_num_format_for_header(header)

    def _write_run_metadata_rows(self, ws):
        super()._write_run_metadata_rows(ws)
        rows = (
            (2, "Pin integration per sample (s)", self.poll_interval_s),
            (3, "Samples per load point", self._sample_count()),
            (4, "Pin integration (s)", self._effective_integration_s()),
            (5, "Est. poll window (s)", round(self._estimated_poll_point_duration_s(), 1)),
            (6, "Fan cooling (HMP CH2)", "Yes" if self.enable_fan_cooling else "No"),
        )
        for row, label, value in rows:
            ws[f"D{row}"] = label
            ws[f"E{row}"] = value
            ws[f"D{row}"].font = CellFont(bold=True, size=10, color="1F497D")
            ws[f"E{row}"].font = CellFont(size=10)

    def _write_excel_section_header_row(self, ws, row: int):
        super()._write_excel_section_header_row(ws, row)
        from openpyxl.utils import get_column_letter
        first_col = self._poll_first_excel_col()
        last_col = self.scope_excel_col() - 1
        first_letter = get_column_letter(first_col)
        last_letter = get_column_letter(last_col)
        ws.merge_cells(f"{first_letter}{row}:{last_letter}{row}")
        cell = ws[f"{first_letter}{row}"]
        cell.value = "Timed Poll"
        self._style_excel_section_title_cell(cell)

    def setup_data_file(self):
        super().setup_data_file()
        desired = (
            f"{self.short_title}_{self.coupling}_{round(self.vout_V, 3):g}V_"
            f"{round(self.i_max_A, 3):g}A")
        if len(desired) > 31:
            desired = desired[:31]
        if self.sheet_name == desired:
            return
        ws = self.wb[self.sheet_name]
        ws.title = desired
        self.sheet_name = desired
        self.wb.save(self.data_file_path)
        self.wb.close()
        self.wb = openpyxl.load_workbook(self.data_file_path)
        self.ws = self.wb[self.sheet_name]

    def test_loop(self):
        try:
            super().test_loop()
        except (TestStopped, TestSkipped):
            self._save_collected_rows()
            raise

    def _save_collected_rows(self):
        """Write samples already taken when the operator stops or skips."""
        try:
            df = getattr(self, 'output_dataframe', None)
            if df is not None and len(df) > 1 and getattr(self, 'data_file_path', None):
                self._export_line_block()
            if getattr(self, 'data_file_path', None):
                self._apply_main_sheet_number_formats()
            if getattr(self, 'embed_images', False) and getattr(self, 'captured_records', None):
                self._link_capture_main_sheet_rows()
                self.embed_waveforms_in_excel()
            path = getattr(self, 'data_file_path', None)
            if path:
                print(f"[POLL] Partial results saved to {path}")
        except Exception as exc:
            print(f"[Warning] Could not save partial poll workbook: {exc}")

    def _note_partial_capture(self, records_before, rows_at_start):
        """One scope PNG per poll sample; linking is 1:1 (no n_samples blob on first capture)."""
        return

    def _raise_if_abort(self, records_before, rows_at_start):
        stop = test_control_flags.get('StopTest', False)
        skip = test_control_flags.get('SkipTest', False)
        if not stop and not skip:
            return
        self._note_partial_capture(records_before, rows_at_start)
        if stop:
            raise TestStopped
        raise TestSkipped

    def _wait_or_stop(self, seconds, records_before, rows_at_start):
        """Sleep in short slices so Stop is honored during the poll interval."""
        remaining = float(seconds)
        while remaining > 1e-9:
            self._raise_if_abort(records_before, rows_at_start)
            chunk = min(0.25, remaining)
            sleep(chunk)
            remaining -= chunk

    def _wait_for_scheduled_poll_sample(
            self, interval_s, records_before, rows_at_start):
        """Idle wait after previous Time Sampled before the next integration starts."""
        if self.poll_prev_sampled_dt is None:
            return
        gap = max(0.0, float(interval_s))
        if gap <= 0:
            return
        target_start = (
            self.poll_prev_sampled_dt + datetime.timedelta(seconds=gap)
        )
        while datetime.datetime.now() < target_start:
            self._raise_if_abort(records_before, rows_at_start)
            remaining = (target_start - datetime.datetime.now()).total_seconds()
            self._wait_or_stop(min(0.25, max(remaining, 0.0)),
                               records_before, rows_at_start)

    def prepare_source_power_meter_for_poll_input(self):
        """Once per load point: set source timer from Pin integration per sample (param 10)."""
        pm = self.power_meter_source
        if pm is not None and hasattr(pm, '_stop_and_reset_integration'):
            pm._stop_and_reset_integration()
        self.ensure_source_power_meter_auto_range(settle_s=0)
        if pm is None:
            return
        integ = self._effective_integration_s()
        if integ >= 1.0:
            sec = max(1, int(round(integ)))
            pm.integration_settings(mode='NORMAL', timer_s=sec, reset=True)
            try:
                pm.enable_averaging(False)
            except Exception:
                pass
        else:
            try:
                pm.integration_settings(mode='NORMAL', timer_s=0, reset=True)
            except Exception:
                pass
            try:
                pm.enable_averaging(True)
            except Exception:
                pass

    def _poll_per_sample_overhead_s(self) -> float:
        return max(
            float(getattr(self, '_poll_repeat_read_duration_s', self.POLL_READ_DURATION_GUESS_S)),
            self.POLL_READ_DURATION_GUESS_S,
        ) + self.POLL_SCOPE_CAPTURE_GUESS_S

    def _estimated_poll_point_duration_s(self) -> float:
        n = self._sample_count()
        if n <= 0:
            return 0.0
        integ = max(self._effective_integration_s(), 0.0)
        per = integ + self._poll_per_sample_overhead_s()
        return n * per

    def _integrate_source_pin_before_read(self, records_before, rows_at_start):
        """Integrate Pin for T_int before meter read; honors Stop in 0.25 s slices."""
        self.poll_integ_start_dt = None
        self.poll_integ_end_dt = None
        integ = self._effective_integration_s()
        if integ < 1.0:
            return
        pm = self.power_meter_source
        if pm is None:
            return
        self._raise_if_abort(records_before, rows_at_start)
        self.poll_integ_start_dt = datetime.datetime.now()
        pm.start_integration()
        deadline = time.time() + float(integ)
        pmo = getattr(pm, 'power_meter_object', None)
        while time.time() < deadline:
            self._raise_if_abort(records_before, rows_at_start)
            if pmo is not None and hasattr(pmo, 'check_integration_status'):
                try:
                    pmo.check_integration_status()
                except Exception:
                    pass
                if getattr(pmo, '_integration_done', False):
                    self.poll_integ_end_dt = datetime.datetime.now()
                    return
            remaining = deadline - time.time()
            if remaining <= 0:
                break
            sleep(min(0.25, remaining))
        self.poll_integ_end_dt = datetime.datetime.now()

    def _measure_poll_readings(self, i2_A):
        """PM reads for poll samples without re-sending AUTO range every tick."""
        pin_done = self._effective_integration_s() >= 1.0
        self.test_data.gather_data(
            coupling=self.coupling,
            usb_pd=self.usbpd_test,
            stable_source_pin=True,
            source_pin_integration_done=pin_done,
        )
        if self.power_meter_load_2 or self.electronic_load_2:
            self.test_data_2.gather_data_load(
                integrate=False, force_use_eload_data=self.use_eload_data)
            self.test_data_2.vreg_pct = (
                ((self.test_data_2.vout_V - self.vout_2_V) / self.vout_2_V) * 100
                if self.vout_2_V else 0)
        else:
            self.test_data_2.vout_V = self.vout_2_V
            self.test_data_2.iout_A = i2_A
            self.test_data_2.pout_W = self.vout_2_V * i2_A
            self.test_data_2.vreg_pct = 0.0

        self.po_total_W = self.test_data.pout_W + self.test_data_2.pout_W
        self.eff_total_pct = (
            (self.po_total_W / self.test_data.pin_W * 100)
            if self.test_data.pin_W > 0 else 0)

    def _extra_acquire_time_s(self) -> float:
        return self._estimated_poll_point_duration_s()

    def _append_measurement_row(self, scope_row):
        super()._append_measurement_row(scope_row)
        if self.poll_row_sampled_dt is not None:
            self.poll_prev_sampled_dt = self.poll_row_sampled_dt

    def _scope_row_from_measurements(self, meas: dict, trig_level, img_path: str = "") -> list:
        try:
            level = round(float(trig_level), 2)
        except (TypeError, ValueError):
            level = 0.0
        row = [f"CH{self.trigger_channel}", level]
        for label in self.scope_measurement_labels:
            val = meas.get(label)
            row.append(val if val is not None else "N/A")
        if not self.embed_images:
            row.append(img_path or "")
        return row

    def _acquire_operating_point(self, vin_set, i1_A, i2_A):
        n = self._sample_count()
        t_int = float(self.poll_interval_s or 0)
        records_before = len(self.captured_records)
        rows_at_start = len(self.output_dataframe)
        self.poll_prev_sampled_dt = None
        self._poll_repeat_read_duration_s = self.POLL_READ_DURATION_GUESS_S
        self.prepare_load_power_meters_for_poll_output()
        self.prepare_source_power_meter_for_poll_input()
        try:
            for sample in range(n):
                self._raise_if_abort(records_before, rows_at_start)
                self.poll_sample_index = sample + 1
                self.poll_elapsed_s = sample * t_int
                if sample == 0:
                    self.poll_point_start_dt = datetime.datetime.now()
                    self.poll_point_end_dt = (
                        self.poll_point_start_dt
                        + datetime.timedelta(seconds=self._estimated_poll_point_duration_s())
                    )
                self.update_status_log(
                    f"Poll sample {self.poll_sample_index}/{n} "
                    f"(Pin integrate {self._effective_integration_s():g} s)",
                    upcoming="Read meters")
                self._integrate_source_pin_before_read(
                    records_before, rows_at_start)
                read_t0 = datetime.datetime.now()
                self.update_status_log(
                    f"Poll sample {self.poll_sample_index}/{n}",
                    upcoming="Capture Waveform")
                self._measure_poll_readings(i2_A)
                self._poll_waveform_filename_tag = f"_S{self.poll_sample_index}"
                scope_row = self.capture_waveform_at_point(vin_set, i1_A, i2_A)
                read_t1 = datetime.datetime.now()
                self._poll_repeat_read_duration_s = max(
                    0.05,
                    (read_t1 - read_t0).total_seconds(),
                )
                if len(scope_row) > 1 and scope_row[1] == "N/A":
                    self.poll_row_sampled_dt = datetime.datetime.now()
                    self._append_measurement_row(scope_row)
                    return
                self.poll_row_sampled_dt = datetime.datetime.now()
                self._append_measurement_row(scope_row)
        except (TestStopped, TestSkipped):
            self._note_partial_capture(records_before, rows_at_start)
            raise

    def _link_capture_main_sheet_rows(self):
        """One embedded PNG per poll sample → one main sheet row per capture record."""
        super()._link_capture_main_sheet_rows()

    def update_test_list_text(self):
        super().update_test_list_text()
        fan = (
            f"HMP CH{self.HMP_FAN_CHANNEL} fan {self.FAN_COOLING_VOLTAGE_V:g} V / "
            f"{self.FAN_COOLING_CURRENT_A:g} A."
            if self.enable_fan_cooling
            else ""
        )
        note = (
            f"Poll: {self._sample_count()} samples/load point, "
            f"{self.poll_interval_s:g} s Pin integration per sample "
            f"(integrate, read, scope, row). "
            f"{fan} "
            "Stop keeps rows already saved to Excel."
        ).replace("  ", " ")
        if note not in self.test_list_text:
            self.test_list_text += "\n" + note


SteadyStatePollWaveformTest_2Port.tc_default = SteadyStatePollWaveformTest_2Port._build_tc_default()
