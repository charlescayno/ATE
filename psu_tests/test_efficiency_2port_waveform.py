import os
import math

import openpyxl
from openpyxl.styles import Font as CellFont, Alignment
from openpyxl.utils import get_column_letter
try:
    from openpyxl.drawing.image import Image as OpenpyxlImage
except ImportError:
    OpenpyxlImage = None

from psu_tests.test_object_imports import *
from psu_tests.test_efficiency_2port import EfficiencyTest_2Port


class _CapturePromptHelper(QObject):
    show_prompt_sig = Signal(str, str)

    def __init__(self):
        super().__init__()
        self.result = "CAPTURE"
        from PySide2.QtCore import Qt
        self.show_prompt_sig.connect(self.show_dialog, Qt.BlockingQueuedConnection)

    @Slot(str, str)
    def show_dialog(self, title, message):
        from PySide2.QtWidgets import QMessageBox
        from PySide2.QtCore import Qt
        msg_box = QMessageBox()
        msg_box.setWindowFlags(msg_box.windowFlags() | Qt.WindowStaysOnTopHint)
        msg_box.setWindowTitle(title)
        msg_box.setText(message)
        msg_box.setIcon(QMessageBox.Question)

        btn_capture = msg_box.addButton("Capture & Proceed", QMessageBox.AcceptRole)
        btn_recapture = msg_box.addButton("Recapture / Retrigger", QMessageBox.ActionRole)
        btn_skip = msg_box.addButton("Skip Point", QMessageBox.DestructiveRole)
        btn_stop = msg_box.addButton("Stop Test", QMessageBox.RejectRole)
        msg_box.setDefaultButton(btn_capture)

        msg_box.exec_()
        clicked = msg_box.clickedButton()
        if clicked == btn_recapture:
            self.result = "RECAPTURE"
        elif clicked == btn_skip:
            self.result = "SKIP"
        elif clicked == btn_stop:
            self.result = "STOP"
        else:
            self.result = "CAPTURE"


class EfficiencyWaveformTest_2Port(EfficiencyTest_2Port):
    """
    Steady-State Waveform 2 Port: dual-output efficiency with scope capture at each point.
    Main sheet holds full measurements; the Waveforms sheet shows load captions and images.
    """

    LEGACY_TITLE = "Efficiency 2 Port (Waveform)"
    title = "Steady-State Waveform 2 Port"
    short_title = "SS_WF_2P"
    EMBED_IMAGE_MAX_WIDTH_PX = 1200
    DISCHARGE_CC_A = 1.0
    STOP_DISCHARGE_MIN_PULSES = 2
    use_average_efficiency_limits = False
    reapply_scope_before_capture = True
    excel_section_header_row = 14
    excel_first_data_row = 15

    ui_definitions = General_UI_Definitions()
    ui_definitions.stack_page_1 = StackWidget1Pages.LineVoltageRange
    ui_definitions.stack_page_2 = StackWidget2Pages.LoadCurrentRange
    ui_definitions.stack_page_3 = StackWidget3Pages.I2C_Options
    ui_definitions.usb_pd_device_toggle_visible = True
    ui_definitions.nominal_iout_visible = True
    ui_definitions.nominal_vout_visible = True
    ui_definitions.nominal_vout_enable = True
    ui_definitions.load_type_visible = True
    ui_definitions.load_direction_cbx_enabled = False
    ui_definitions.load_range_selection_enabled = True
    ui_definitions.measure_ripple_visible = False

    ui_update = General_UI_Update_Definitions()
    ui_update.line_settings_update = True
    ui_update.load_settings_update = True
    ui_update.soaktime_settings_update = True
    ui_update.cvcc_settings_update = False
    ui_update.line_ramp_settings_update = False
    ui_update.nominal_output_settings_update = True
    ui_update.usbpd_options_update = True
    ui_update.tracking_pdo_request_update = True
    ui_update.measure_ripple_update = False
    ui_update.load_direction_update = True
    ui_update.eload_type_update = True
    ui_update.use_eload_data_update = True
    ui_update.i2c_params_update = True

    # Param indices 1-4 and cbx 1 must match EfficiencyTest_2Port (Add Test dual-load wiring relies on them).
    i2c_ui_definitions = I2C_UI_Definitions()
    i2c_ui_definitions.add_lineedit("Nominal Vout 2 (V)", 1)
    i2c_ui_definitions.add_lineedit("Nominal Iout 2 (A)", 2)
    i2c_ui_definitions.add_lineedit("Fixed Aux Load (A)", 3)
    i2c_ui_definitions.add_lineedit("Port 2 Loads (%)", 4)
    i2c_ui_definitions.add_lineedit("Trigger Delta", 5)
    i2c_ui_definitions.add_lineedit("Settle Time (s)", 6)
    i2c_ui_definitions.add_lineedit("Discharge Pulses", 7)
    i2c_ui_definitions.add_lineedit("Ambient Temp (°C)", 8)
    i2c_ui_definitions.add_lineedit("Post-Capture Delay (s)", 9)
    i2c_ui_definitions.add_lineedit("Trigger Channel (1-4)", 10)
    i2c_ui_definitions.add_cbx("Dual Load Mode", [
        "Option A: Proportional Sync (Both Ports 100% -> 10%)",
        "Option B: Cross-Reg Matrix (Port 2 at Specified % Loads)",
        "Option C: Fixed Aux / Swept Main (Port 2 Fixed Load)"
    ], 1)
    i2c_ui_definitions.add_cbx("Embed in Excel", ['Yes', 'No'], 2)
    i2c_ui_definitions.add_cbx("User Prompts", [
        'Pre-Test & Per-Capture',
        'Pre-Test Only',
        'Per-Capture Only',
        'None (Fully Automated)'
    ], 3)
    i2c_ui_definitions.add_cbx("Auto Find Trigger", ['No', 'Yes'], 4)

    tc_default = TestConditions(
        name=title,
        nominal_output_voltage_V=48.0,
        nominal_load_current_A=3.0,
        max_load_current_A=3.0,
        line_range=LineSettings.TST_326_HVDC,
        load_range=LoadSettings.LOAD_SINGLE_VALUE_100_PCT,
        soak_time=SoaktimeSettings.TST_326,
        general_options=GeneralOptions(
            measure_ripple=False,
            use_eload_data=False,
            eload_type='CC',
            load_direction='Downward',
            coupling='DC',
        ),
        usbpd_options=USBPDOptions(),
        line_ramp_settings=LineRamp(),
        i2c_test_parameters=I2CTestParameters(
            params=[20.0, 0.1, 0.1, 0.0, 3.0, 1.0, 2.0, 25.0, 1.0, 1.0],
            cbx_params=[
                "Option C: Fixed Aux / Swept Main (Port 2 Fixed Load)",
                "Yes",
                "Per-Capture Only",
                "Yes",
            ]),
        unit_id="00",
        ambient_temp=25.0,
        scope_channels={
            1: {
                'enabled': True,
                'name': 'VDS_SR',
                'measurements': ['MAXimum'],
                'derating': 150.0,
                'cursor': {'enabled': False, 'type': 'VERTical'},
            },
            2: {
                'enabled': False,
                'name': 'VDS_PRI',
                'measurements': [],
                'derating': None,
                'cursor': {},
            },
            3: {'enabled': False, 'name': '', 'measurements': [], 'derating': None, 'cursor': {}},
            4: {'enabled': False, 'name': '', 'measurements': [], 'derating': None, 'cursor': {}},
        })

    def __init__(self, test_item):
        super().__init__(test_item)
        self.captured_records = []

    def with_waveform_capture(self):
        return True

    def run(self):
        if self.parent.run_settings['debug']:
            debugpy.debug_this_thread()
        global test_control_flags
        try:
            self.status_update.emit(TestStatus.IN_PROGRESS)
            self.setup_equipment()
            self.prepare_test_conditions()
            self.setup_data_file()
            self.test_loop()
        except TestStopped:
            self.safe_output_discharge(reason="stopped")
            self.cleanup_usbpd_pps_operation()
            self.input_supply_eload_discharge_sequence()
            print("Test Stopped")
            self.status_update.emit(TestStatus.STOPPED)
        except TestSkipped:
            self.cleanup_usbpd_pps_operation()
            self.input_supply_eload_discharge_sequence()
            print("Test Skipped")
            self.status_update.emit(TestStatus.SKIPPED)
        except Exception:
            print(traceback.format_exc())
            with open('error_log.txt', 'a') as f:
                f.write(traceback.format_exc() + '\n')
            self.cleanup_usbpd_pps_operation()
            if self.usbpd_test and self.usbpd_sink:
                self.usbpd_sink.usb_pd_initialize()
            self.input_supply_eload_discharge_sequence()
            print("Test Failed")
            self.status_update.emit(TestStatus.FAILED)
        else:
            self.cleanup_usbpd_pps_operation()
            self.input_supply_eload_discharge_sequence()
            self.estimated_time_s = 0
            self.status_update.emit(TestStatus.COMPLETE)

    def safe_output_discharge(self, reason: str = ""):
        """CCH discharge on assigned e-loads; used when the operator stops the test."""
        pulses = max(int(getattr(self, 'discharge_pulses', 0) or 0), self.STOP_DISCHARGE_MIN_PULSES)
        print(f"[DISCHARGE] safe output ({reason}): {pulses} pulse(s) at {self.DISCHARGE_CC_A:g} A CCH")
        self.discharge_output(pulses)

    ###########################################################################
    #                         Test item / parameters                          #
    ###########################################################################
    def unpack_test_item(self):
        super().unpack_test_item()
        # Disable the parent's optional scope sub-table; this test writes scope data inline.
        self.measure_ripple = False

        tc = self.test_conditions
        self.unit_id = str(getattr(tc, 'unit_id', '00') or '00')

        self.scope_channels = getattr(tc, 'scope_channels', None) or {
            1: {'enabled': True, 'name': 'Primary Vds'},
            2: {'enabled': True, 'name': 'Ids'},
            3: {'enabled': False, 'name': ''},
            4: {'enabled': False, 'name': ''}
        }
        self.scope_channels = {int(k): v for k, v in self.scope_channels.items()}
        active = []
        for ch_num, ch_data in sorted(self.scope_channels.items()):
            if ch_data.get('enabled', False):
                name = str(ch_data.get('name', '')).strip()
                active.append(f"CH{ch_num}: {name}" if name else f"CH{ch_num}")
        self.probe_setup = ", ".join(active) if active else "None"

        i2c = getattr(tc, 'i2c_test_parameters', None)
        params = list(getattr(i2c, 'param', []) or [])
        cbx = list(getattr(i2c, 'cbx_param', []) or [])

        def _param(index, default, cast=float, minimum=None):
            try:
                if index < len(params) and str(params[index]).strip() != '':
                    val = cast(float(params[index]))
                    if minimum is not None and val < minimum:
                        return default
                    return val
            except (ValueError, TypeError):
                pass
            return default

        self.trigger_delta = _param(4, 0.5, minimum=1e-9)
        self.settle_time = _param(5, 1.0, minimum=0)
        self.discharge_pulses = _param(6, 0, cast=int, minimum=0)
        self.ambient_temp = _param(7, float(getattr(tc, 'ambient_temp', 25.0) or 25.0))
        self.post_capture_delay = _param(8, 1.0, minimum=0)
        trig = _param(9, 1, cast=int)
        self.trigger_channel = trig if trig in (1, 2, 3, 4) else 1

        self.embed_images = (str(cbx[1]).strip() != 'No') if len(cbx) > 1 and cbx[1] else True
        self.user_prompts = str(cbx[2]).strip() if len(cbx) > 2 and cbx[2] else 'Pre-Test & Per-Capture'
        self.auto_find_trigger = (str(cbx[3]).strip() == 'Yes') if len(cbx) > 3 and cbx[3] else False
        self.prompt_before_start = 'Pre-Test' in self.user_prompts
        self.prompt_before_capture = 'Per-Capture' in self.user_prompts

    def setup_equipment(self):
        super().setup_equipment()
        if self.oscilloscope is None and hasattr(self.equipment, 'check_scope_availability'):
            try:
                self.equipment.check_scope_availability(get_default=True)
                self.oscilloscope = self.equipment.oscilloscope
            except Exception:
                pass
        if self.oscilloscope is None:
            raise ConnectionError(
                "No Oscilloscope Connected or Accessible. Please verify the oscilloscope "
                "in Equipment Setup.")

    ###########################################################################
    #                           Data file / headers                           #
    ###########################################################################
    def define_scope_header(self):
        labels = []
        for ch_idx, ch_data in sorted(self.scope_channels.items()):
            if not ch_data.get('enabled', False):
                continue
            ch_name = ch_data.get('name') or f"CH{ch_idx}"
            for meas in (ch_data.get('measurements') or ['MAXimum']):
                lbl = self._normalize_meas_label(meas)
                col = f"{ch_name} {lbl}"
                if col not in labels:
                    labels.append(col)
                if lbl == "Max" and ch_data.get('derating') is not None:
                    der = f"{ch_name} Derating (%)"
                    if der not in labels:
                        labels.append(der)
            if (ch_data.get('cursor') or {}).get('enabled'):
                for suffix in ("dX", "Fsw (kHz)", "dY"):
                    col = f"{ch_name} {suffix}"
                    if col not in labels:
                        labels.append(col)
        if not labels:
            labels = ['Vds Max (V)', 'Ids Max (A)']
        self.scope_measurement_labels = labels
        scope_cols = ['Trig Ch', 'Trig Level (V)'] + labels
        if not getattr(self, 'embed_images', True):
            scope_cols.append('Waveform File')
        return scope_cols

    def setup_data_file(self):
        self.test_data = TestData()
        self.test_data.vout_nom_V = self.vout_V
        self.test_data.use_eload_data = self.use_eload_data
        self.test_data.source_power_meter = self.power_meter_source
        self.test_data.load_power_meter = self.power_meter_load
        self.test_data.electronic_load = self.electronic_load

        self.test_data_2 = TestData()
        self.test_data_2.vout_nom_V = self.vout_2_V
        self.test_data_2.use_eload_data = self.use_eload_data
        self.test_data_2.load_power_meter = self.power_meter_load_2
        self.test_data_2.electronic_load = self.electronic_load_2

        self.header_list = self.define_data_header()
        self.scope_header_list = self.define_scope_header()
        self.output_dataframe = dataframe_from_headers(self.header_list)
        self.scope_dataframe = dataframe_from_headers(self.scope_header_list)
        if self.use_average_efficiency_limits:
            self.efficiency_header = ['Avg_Eff', 'DOE6 Limit', 'COC5 T2 Limit', 'COC5_T2_10%', 'Pass/Fail']
            self.efficiency_dataframe = dataframe_from_headers(self.efficiency_header)
        else:
            self.efficiency_header = []
            self.efficiency_dataframe = None

        if not os.path.exists(self.output_folder_path):
            os.makedirs(self.output_folder_path, exist_ok=True)
        if not self.waveform_filepath:
            self.waveform_filepath = f"{self.output_folder_path}/waveforms"
        os.makedirs(self.waveform_filepath, exist_ok=True)

        self.data_filename = f'{self.title} Test {self.vout_V:g}V'
        self.data_file_path = f'{self.output_folder_path}/{self.data_filename}.xlsx'
        if not os.path.exists(self.data_file_path):
            wb = openpyxl.Workbook()
            wb.save(self.data_file_path)
            wb.close()

        self.wb = openpyxl.load_workbook(self.data_file_path)
        self.sheet_name = (
            f"SS_WF_2P_{self.coupling}_{round(self.vout_V, 3):g}V_{round(self.i_max_A, 3):g}A")
        if self.sheet_name in self.wb.sheetnames:
            clear_sheet(self.output_folder_path, self.data_filename, self.sheet_name)
            self.wb = openpyxl.load_workbook(self.data_file_path)
            self.ws = self.wb[self.sheet_name]
        else:
            self.ws = self.wb.create_sheet(title=self.sheet_name)
            if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
                del self.wb['Sheet']

        self.prepare_sheet_formatting()
        self.next_excel_data_row = self.excel_first_data_row
        self._excel_vin_block_index = 0
        self.define_output_data_objects()

    def scope_excel_col(self) -> int:
        """1-based column index for the first oscilloscope column."""
        return 25 if self.use_average_efficiency_limits else 20

    def _write_run_metadata_rows(self, ws):
        soak = self.soak_time
        meta = [
            ("Unit ID", self.unit_id),
            ("Ambient Temp (°C)", self.ambient_temp),
            ("Ch1", f"{self.vout_V:g} V / {self.i_max_A:g} A"),
            ("Ch2", f"{self.vout_2_V:g} V / {self.i_max_2_A:g} A"),
            ("Mode", self.dual_load_mode),
            ("Probes", self.probe_setup),
            ("Trigger", f"CH{self.trigger_channel}"),
            ("Soak profile", soak.name),
            ("Initial soak (s)", soak.initial_soak),
            ("Soak per line (s)", soak.soak_per_line),
            ("Soak per load (s)", soak.soak_per_load),
            ("Integration time (s)", soak.integration_time),
        ]
        for row_offset, (label, value) in enumerate(meta):
            row = 2 + row_offset
            ws[f"A{row}"] = label
            ws[f"B{row}"] = value
            ws[f"A{row}"].font = CellFont(bold=True, size=10, color="1F497D")
            ws[f"B{row}"].font = CellFont(size=10)

    def _excel_section_title_font(self):
        return CellFont(bold=True, size=11, color="1F497D")

    def _excel_column_header_font(self):
        return CellFont(bold=True, size=10)

    def _style_excel_section_title_cell(self, cell):
        cell.font = self._excel_section_title_font()
        cell.alignment = Alignment(horizontal='center', vertical='center')

    def _write_excel_section_header_row(self, ws, row: int):
        ws.merge_cells(f"A{row}:B{row}")
        ws[f"A{row}"] = 'Load Condition'
        self._style_excel_section_title_cell(ws[f"A{row}"])
        ws.merge_cells(f"C{row}:D{row}")
        ws[f"C{row}"] = 'Input Set'
        self._style_excel_section_title_cell(ws[f"C{row}"])
        ws.merge_cells(f"E{row}:I{row}")
        ws[f"E{row}"] = 'Input Measurement'
        self._style_excel_section_title_cell(ws[f"E{row}"])
        ws.merge_cells(f"J{row}:M{row}")
        ws[f"J{row}"] = 'Output 1 Measurement'
        self._style_excel_section_title_cell(ws[f"J{row}"])
        ws.merge_cells(f"N{row}:Q{row}")
        ws[f"N{row}"] = 'Output 2 Measurement'
        self._style_excel_section_title_cell(ws[f"N{row}"])
        ws.merge_cells(f"R{row}:S{row}")
        ws[f"R{row}"] = 'Total Power & Efficiency'
        self._style_excel_section_title_cell(ws[f"R{row}"])
        if self.use_average_efficiency_limits:
            ws.merge_cells(f"T{row}:X{row}")
            ws[f"T{row}"] = 'Average Efficiency Limits'
            self._style_excel_section_title_cell(ws[f"T{row}"])
        first_col = self.scope_excel_col()
        last_col = first_col + len(self.scope_header_list) - 1
        ws.merge_cells(
            f"{get_column_letter(first_col)}{row}:{get_column_letter(last_col)}{row}")
        cell = ws[f"{get_column_letter(first_col)}{row}"]
        cell.value = 'Oscilloscope Measurements'
        self._style_excel_section_title_cell(cell)

    def _bold_excel_column_header_row(self, row: int):
        """Bold Ch1 Load, Vin, scope columns, etc. on the dataframe header row."""
        wb = openpyxl.load_workbook(self.data_file_path)
        ws = wb[self.sheet_name]
        last_col = self.scope_excel_col() + len(self.scope_header_list) - 1
        header_font = self._excel_column_header_font()
        center = Alignment(horizontal='center', vertical='center')
        for col in range(1, last_col + 1):
            cell = ws.cell(row=row, column=col)
            if cell.value is None or str(cell.value).strip() == '':
                continue
            cell.font = header_font
            cell.alignment = center
        wb.save(self.data_file_path)
        wb.close()

    def prepare_sheet_formatting(self):
        ws = self.ws
        ws['A1'] = f"{self.title} Test"
        ws['A1'].font = CellFont(size=14, bold=True, color="1F497D")
        self._write_run_metadata_rows(ws)
        self._write_excel_section_header_row(ws, self.excel_section_header_row)
        self.wb.save(self.data_file_path)
        self.wb.close()

    ###########################################################################
    #                               Test loop                                 #
    ###########################################################################
    def create_message_popup(self, title: str, message: str, message_type: MessageType):
        if os.environ.get('PI_ATE_AUTO_CONFIRM') == '1':
            return
        super().create_message_popup(title, message, message_type)

    def update_status_log(self, msg, upcoming=None):
        print(f"[TEST_STEP] {msg}")
        self.current_status_log = msg
        self.upcoming_event_log = upcoming

    def apply_scope_setup(self):
        """Label channels and push custom measurements/cursors to the scope when supported."""
        custom_meas = {}
        for ch_num, ch_data in self.scope_channels.items():
            ch_name = (ch_data.get('name') or f"CH{ch_num}").strip()
            enabled = bool(ch_data.get('enabled'))
            if hasattr(self.oscilloscope, 'apply_channel_from_ui'):
                try:
                    self.oscilloscope.apply_channel_from_ui(ch_num, enabled, ch_name)
                except Exception as e:
                    print(f"[Warning] Failed to apply CH{ch_num} on scope: {e}")
                if not enabled:
                    continue
            elif not enabled:
                if hasattr(self.oscilloscope, 'hide_channel'):
                    try:
                        self.oscilloscope.hide_channel(ch_num)
                    except Exception as e:
                        print(f"[Warning] Failed to hide CH{ch_num}: {e}")
                continue
            else:
                if hasattr(self.oscilloscope, 'show_channel'):
                    try:
                        self.oscilloscope.show_channel(ch_num)
                    except Exception as e:
                        print(f"[Warning] Failed to show CH{ch_num}: {e}")
                if hasattr(self.oscilloscope, 'channel_label'):
                    try:
                        self.oscilloscope.channel_label(channel=ch_num, label=ch_name,
                                                        rel_x_position=50, rel_y_position=50)
                    except Exception as e:
                        print(f"[Warning] Failed to set label '{ch_name}' for CH{ch_num}: {e}")
            if ch_data.get('measurements'):
                custom_meas[ch_num] = ch_data.get('measurements')
            if (ch_data.get('cursor') or {}).get('enabled') and hasattr(self.oscilloscope, 'cursor'):
                try:
                    ctype = ch_data['cursor'].get('type', 'VERTical')
                    ctype_short = ctype[:4].upper() if len(ctype) >= 4 else 'VERT'
                    if ctype_short == 'HORI':
                        ctype_short = 'HOR'
                    self.oscilloscope.cursor(channel=ch_num, cursor_set=ch_num, type=ctype_short)
                except Exception as e:
                    print(f"[Warning] Failed to set cursor on CH{ch_num}: {e}")
        if custom_meas:
            print(f"[SCOPE] apply_scope_setup measurements: {custom_meas}")
        if custom_meas and hasattr(self.oscilloscope, 'set_channel_measurements'):
            try:
                self.oscilloscope.set_channel_measurements(custom_meas)
                sleep(1)
                configured = getattr(self.oscilloscope, 'configured_measurements', None)
                if configured is not None:
                    print(f"[SCOPE] configured_measurements: {configured}")
            except Exception as e:
                print(f"[Warning] Failed to set custom measurements on scope: {e}")
        for ch_num, ch_data in self.scope_channels.items():
            if not (ch_data.get('cursor') or {}).get('enabled'):
                continue
            if not hasattr(self.oscilloscope, 'write'):
                continue
            try:
                state = self.oscilloscope.write(f'CURS{ch_num}:STAT?')
                print(f"[SCOPE] CH{ch_num} cursor STAT readback: {state!r}")
            except Exception as e:
                print(f"[Warning] Cursor readback CH{ch_num}: {e}")

    def test_loop(self):
        soak = self.soak_time

        if self.prompt_before_start:
            prompt_msg = (
                "<b>================== OSCILLOSCOPE SETUP REMINDER ==================</b><br>"
                f"Unit ID: {self.unit_id} | Ambient Temp: {self.ambient_temp:g}\u00b0C<br>"
                f"Ch1: {self.vout_V:g} V / {self.i_max_A:g} A &nbsp; "
                f"Ch2: {self.vout_2_V:g} V / {self.i_max_2_A:g} A &nbsp; Mode: {self.dual_load_mode}<br><br>"
                "1. Load the correct .DFL file into the oscilloscope.<br>"
                f"2. Probe Configuration: {self.probe_setup}<br>"
                f"3. Trigger Channel: CH{self.trigger_channel} (Delta: {self.trigger_delta:g} V)<br>"
                "<b>=================================================================</b><br><br>"
                "Click OK once setup is complete to proceed."
            )
            self.create_message_popup("Oscilloscope Setup Verification", prompt_msg, MessageType.INFO)

        self.apply_scope_setup()

        self.total_time, self.total_steps = self.estimate_remaining(0, 0, vin_delays=True)
        self.status_report(0, 0, vin_delays=True)
        if self.usbpd_test:
            self.input_supply_eload_discharge_sequence()

        self.confirm_discharge_resistor_disengaged()

        for self.vin_index, self.vin_freq in enumerate(self.vin_list):
            vin_set = self.vin_freq[0]
            self.update_status_log(f"Setting Input {vin_set:g} V{self.coupling}", upcoming="Line Soak")
            self.input_supply.set_voltage_with_coupling(voltage=vin_set, coupling=self.coupling)
            if self.coupling == AC_SOURCE_COUPLING.AC:
                self.input_supply.frequency = self.vin_freq[1]
            self.source_vout = vin_set

            self.test_data.vin_set_V = vin_set
            self.test_data.ac_freq_Hz = self.vin_freq[1]

            self.create_new_plot_series()
            self.input_supply.turn_on()

            if self.power_meter_load:
                self.power_meter_load.auto_range_enable()
            if self.power_meter_load_2:
                self.power_meter_load_2.auto_range_enable()
            if self.power_meter_source:
                self.power_meter_source.current_auto_range_enable()
            sleep(2)

            if self.usbpd_test:
                usbpd_init_ok = False
                counter = 0
                while not usbpd_init_ok:
                    usbpd_init_ok = self.usbpd_sink.usb_pd_initialize()
                    sleep(0.2)
                    counter += 1
                    if counter == 20:
                        raise ConnectionError("No USB-PD Source Connected")
                counter = 0
                while self.usbpd_sink.source_cap_count == 0:
                    self.usbpd_sink.get_status(serial_number=self.usbpd_sink.serial_number)
                    counter += 1
                    sleep(0.2)
                    if counter == 20:
                        raise ConnectionError("Unable to Get Source Capabilities")
                sleep(1)
                self.usb_pd_request(vout_V=self.vout_V, iout_A=self.i_rated_A)

            init_i1 = self.load_pairs[0][0] if self.load_pairs else 0
            init_i2 = self.load_pairs[0][1] if self.load_pairs else 0
            self._set_loads(init_i1, init_i2)
            sleep(2)
            self.correct_source_output()

            if self.vin_index == 0:
                self.update_status_log("Initial Soak...")
                soak.do_initial_soak()
            else:
                self.update_status_log("Line Soak...")
                soak.do_soak_per_line()

            for load_index, (i1_A, i2_A, ch1_lbl, ch2_lbl) in enumerate(self.load_pairs):
                if test_control_flags.get('StopTest', False):
                    raise TestStopped
                if test_control_flags.get('SkipTest', False):
                    raise TestSkipped

                self.status_report(self.vin_index, load_index, vin_delays=False)
                self.current_ch1_lbl = ch1_lbl
                self.current_ch2_lbl = ch2_lbl
                self.current_load_index = load_index

                self.update_status_log(f"Setting Loads Ch1 {i1_A:g} A / Ch2 {i2_A:g} A", upcoming="Load Soak")
                self._set_loads(i1_A, i2_A)

                if self.power_meter_load is not None:
                    if i1_A < 0.05:
                        self.power_meter_load.set_current_range(0.05)
                    else:
                        self.power_meter_load.current_auto_range_enable()
                if self.power_meter_load_2 is not None:
                    if i2_A < 0.05:
                        self.power_meter_load_2.set_current_range(0.05)
                    else:
                        self.power_meter_load_2.current_auto_range_enable()

                sleep(2)
                self.correct_source_output()
                self.update_status_log("Load Soak...", upcoming="Efficiency Measurement")
                soak.do_soak_per_load()
                sleep(self.settle_time)

                self.update_status_log("Measuring efficiency...", upcoming="Capture Waveform")
                self._measure_efficiency(i2_A)

                scope_row = self.capture_waveform_at_point(vin_set, i1_A, i2_A)

                self.output_dataframe.loc[len(self.output_dataframe)] = self.process_data_row()
                self.scope_dataframe.loc[len(self.scope_dataframe)] = scope_row
                self.current_scope_row = scope_row
                self.update_output_data()

                if self.power_meter_load:
                    self.power_meter_load.auto_range_enable()
                if self.power_meter_load_2:
                    self.power_meter_load_2.auto_range_enable()
                if self.power_meter_source:
                    self.power_meter_source.current_auto_range_enable()

                sleep(self.post_capture_delay)

            self.test_data_table.add_blank_row()
            self._export_line_block()

            if self.usbpd_test:
                self.input_supply_eload_discharge_sequence()
            sleep(3)

        if self.discharge_pulses > 0:
            self.discharge_output(self.discharge_pulses)
        self._apply_main_sheet_number_formats()
        if self.embed_images:
            self._link_capture_main_sheet_rows()
            self.embed_waveforms_in_excel()

    def _set_loads(self, i1_A, i2_A):
        self.electronic_load.set_load(self.vout_V, i1_A, self.eload_type)
        if i1_A > 0:
            self.electronic_load.turn_on()
        else:
            self.electronic_load.turn_off()
        if self.electronic_load_2 is not None:
            self.electronic_load_2.set_load(self.vout_2_V, i2_A, self.eload_type)
            if i2_A > 0:
                self.electronic_load_2.turn_on()
            else:
                self.electronic_load_2.turn_off()

    def _measure_efficiency(self, i2_A):
        if self.power_meter_load and getattr(self.power_meter_load, '_current_auto_range_status', False):
            self.power_meter_load.auto_range_enable(False)
            sleep(0.5)
        if self.power_meter_load_2 and getattr(self.power_meter_load_2, '_current_auto_range_status', False):
            self.power_meter_load_2.auto_range_enable(False)
            sleep(0.5)
        if self.power_meter_source and getattr(self.power_meter_source, '_current_auto_range_status', False):
            self.power_meter_source.current_auto_range_enable(False)
            sleep(0.5)

        self.test_data.gather_data(coupling=self.coupling, usb_pd=self.usbpd_test)
        if self.power_meter_load_2 or self.electronic_load_2:
            self.test_data_2.gather_data_load(integrate=False, force_use_eload_data=self.use_eload_data)
            self.test_data_2.vreg_pct = (
                ((self.test_data_2.vout_V - self.vout_2_V) / self.vout_2_V) * 100 if self.vout_2_V else 0)
        else:
            self.test_data_2.vout_V = self.vout_2_V
            self.test_data_2.iout_A = i2_A
            self.test_data_2.pout_W = self.vout_2_V * i2_A
            self.test_data_2.vreg_pct = 0.0

        self.po_total_W = self.test_data.pout_W + self.test_data_2.pout_W
        self.eff_total_pct = (self.po_total_W / self.test_data.pin_W * 100) if self.test_data.pin_W > 0 else 0

    def _export_dataframe_slice(self, df, anchor: str):
        if df is None or len(df) == 0:
            return 0
        first_block = self._excel_vin_block_index == 0
        export_df = df if first_block else df.iloc[1:]
        if len(export_df) == 0:
            return 0
        export_to_excel(export_df, self.output_folder_path,
                        self.data_filename, self.sheet_name, anchor)
        return len(export_df)

    def _export_line_block(self):
        anchor_row = self.next_excel_data_row
        scope_col = get_column_letter(self.scope_excel_col())
        first_block = self._excel_vin_block_index == 0
        rows_written = self._export_dataframe_slice(
            self.output_dataframe, f"A{anchor_row}")
        if self.use_average_efficiency_limits and self.efficiency_dataframe is not None:
            self._export_dataframe_slice(
                self.efficiency_dataframe, f"T{anchor_row}")
        self._export_dataframe_slice(
            self.scope_dataframe, f"{scope_col}{anchor_row}")
        if first_block:
            self._bold_excel_column_header_row(anchor_row)
        self.next_excel_data_row = anchor_row + rows_written
        self._excel_vin_block_index += 1
        if rows_written > 0:
            start_data = anchor_row + 1 if first_block else anchor_row
            end_row = anchor_row + rows_written - 1
            if end_row >= start_data:
                self._format_main_sheet_data_rows(start_data, end_row)

        self.output_dataframe = dataframe_from_headers(self.header_list)
        if self.efficiency_dataframe is not None:
            self.efficiency_dataframe = dataframe_from_headers(self.efficiency_header)
        self.scope_dataframe = dataframe_from_headers(self.scope_header_list)

    ###########################################################################
    #                          Oscilloscope capture                           #
    ###########################################################################
    def capture_waveform_at_point(self, vin_set, i1_A, i2_A) -> list:
        """Trigger, prompt, and screenshot at the current point.
        Returns the scope row aligned with self.scope_header_list."""
        n_meas = len(self.scope_measurement_labels)
        while True:
            if self.auto_find_trigger:
                self.update_status_log("Sweeping Trigger Level...", upcoming="Capture Waveform")
                trig_level = self.find_trigger(channel=self.trigger_channel, trigger_delta=self.trigger_delta)
            else:
                trig_level = self._read_trigger_level()

            self.update_status_log("Capturing waveform...", upcoming="Save Screenshot")
            if self.reapply_scope_before_capture:
                self.apply_scope_setup()
            self.oscilloscope.run_single()
            sleep(1)
            self.oscilloscope.stop()
            meas = self.query_active_scope_measurements()

            if self.prompt_before_capture:
                meas_str = "".join(f"{k}: {v}<br>" for k, v in meas.items()) or "No active measurements found.<br>"
                message = (
                    f"<b>Operating Point:</b> {vin_set:g} V{self.coupling} | "
                    f"Ch1 {i1_A:g} A ({self.current_ch1_lbl}) | Ch2 {i2_A:g} A ({self.current_ch2_lbl})<br>"
                    f"<b>Efficiency:</b> {self.eff_total_pct:.2f} % &nbsp; "
                    f"<b>Pin:</b> {self.test_data.pin_W:.3f} W &nbsp; <b>Po_Total:</b> {self.po_total_W:.3f} W<br>"
                    f"<b>Unit ID:</b> {self.unit_id}<br>"
                    f"<b>Trigger Level:</b> CH{self.trigger_channel} @ {trig_level:.2f} V<br><br>"
                    f"<b>Live Measurements:</b><br>{meas_str}<br>"
                    "Review the waveform on the oscilloscope screen, then choose an action:"
                )
                action = self.request_capture_action(
                    f"Capture Waveform ({vin_set:g} V, Ch1 {self.current_ch1_lbl} / Ch2 {self.current_ch2_lbl})",
                    message)
                if action == "RECAPTURE":
                    print(f"[RECAPTURE] Re-trigger at {vin_set} V, Ch1 {i1_A} A, Ch2 {i2_A} A")
                    continue
                if action == "STOP":
                    raise TestStopped
                if action == "SKIP":
                    print(f"[SKIP] Capture skipped at {vin_set} V, Ch1 {i1_A} A, Ch2 {i2_A} A")
                    self.oscilloscope.run()
                    row = [f"CH{self.trigger_channel}", "N/A"] + ["N/A"] * n_meas
                    if not self.embed_images:
                        row.append("")
                    return row

            prefix = f"{self.unit_id}_" if self.unit_id else ""
            img_filename = f"{prefix}{vin_set:g}V{self.coupling}_Ch1_{i1_A:g}A_Ch2_{i2_A:g}A.png"
            img_path = os.path.join(self.waveform_filepath, img_filename)
            try:
                self.oscilloscope.get_screenshot(img_filename, self.waveform_filepath)
            except Exception as e:
                print(f"[Warning] Failed to capture oscilloscope screenshot {img_filename}: {e}")
            self.oscilloscope.run()

            self.captured_records.append({
                'vin': vin_set, 'i1': i1_A, 'i2': i2_A,
                'ch1_lbl': self.current_ch1_lbl, 'ch2_lbl': self.current_ch2_lbl,
                'eff': self.eff_total_pct, 'path': img_path, 'measurements': meas,
            })

            row = [f"CH{self.trigger_channel}", round(trig_level, 2)]
            for label in self.scope_measurement_labels:
                val = meas.get(label)
                row.append(val if val is not None else "N/A")
            if not self.embed_images:
                row.append(img_path)
            return row

    def _read_trigger_level(self) -> float:
        try:
            return float(self.oscilloscope.device.query(f'TRIG1:LEV{self.trigger_channel}?'))
        except Exception:
            return 0.0

    def request_capture_action(self, title: str, message: str) -> str:
        if os.environ.get('PI_ATE_AUTO_CONFIRM') == '1':
            return "CAPTURE"
        try:
            from PySide2.QtWidgets import QApplication
            app = QApplication.instance()
            if app:
                helper = _CapturePromptHelper()
                helper.moveToThread(app.thread())
                helper.show_prompt_sig.emit(title, message)
                return helper.result
        except Exception as e:
            print(f"Error displaying capture prompt: {e}")
        return "CAPTURE"

    def find_trigger(self, channel: int = 1, trigger_delta: float = 0.5) -> float:
        """Raise the trigger level until it stops triggering, then back off one step."""
        sc = self.oscilloscope
        sc.run_single()
        if hasattr(sc, 'trigger_mode'):
            sc.trigger_mode('AUTO')
        sleep(1)
        if hasattr(sc, 'trigger_mode'):
            sc.trigger_mode('NORM')
        try:
            labels, values = sc.get_measure(channel)
            try:
                max_value = float(values[0]) if values and values[0] is not None else 10.0
                if math.isnan(max_value):
                    max_value = 10.0
            except (ValueError, TypeError):
                max_value = 10.0

            trigger_level = float(f"{max_value:.4f}")
            sc.edge_trigger(channel, trigger_level, 'POS')
            sc.run_single()
            sleep(1)
            trigger_status = sc.trigger_status()

            steps = 0
            while trigger_status == 1 and steps < 30:
                steps += 1
                if test_control_flags.get('StopTest', False):
                    raise TestStopped
                if test_control_flags.get('SkipTest', False):
                    raise TestSkipped
                trigger_level += trigger_delta
                sc.edge_trigger(channel, trigger_level, 'POS')
                sc.run_single()
                sleep(0.5)
                trigger_status = sc.trigger_status()

            trigger_level -= trigger_delta
            sc.edge_trigger(channel, trigger_level, 'POS')
            sleep(6)
            return trigger_level
        except (TestStopped, TestSkipped):
            raise
        except Exception as e:
            print(f"[Warning] find_trigger error on CH{channel}: {e}")
            return 10.0

    @staticmethod
    def _normalize_meas_label(raw: str) -> str:
        up = raw.upper()
        if "MAX" in up:
            return "Max"
        if "MIN" in up:
            return "Min"
        if "MEAN" in up:
            return "Mean"
        if "RMS" in up:
            return "RMS"
        if "PDEL" in up or "PKPK" in up:
            return "Pk-Pk"
        return raw.title()

    def _channel_name(self, ch_idx):
        ch_info = self.scope_channels.get(ch_idx, {})
        return ch_info.get('name') or f"CH{ch_idx}"

    def query_active_scope_measurements(self) -> dict:
        meas_dict = {}
        try:
            for item in (self.oscilloscope.get_measure_all() or []):
                labels = item.get('labels')
                values = item.get('values')
                ch_idx = item.get('channel', 1)
                if not labels or not values:
                    continue
                ch_info = self.scope_channels.get(ch_idx, {})
                for lbl_raw, val in zip(labels, values):
                    if val is None or (isinstance(val, float) and math.isnan(val)):
                        continue
                    lbl = self._normalize_meas_label(lbl_raw)
                    meas_val = round(float(val), 2)
                    meas_dict[f"{self._channel_name(ch_idx)} {lbl}"] = meas_val
                    if lbl == "Max" and ch_info.get('derating') is not None:
                        limit = ch_info.get('derating')
                        meas_dict[f"{self._channel_name(ch_idx)} Derating (%)"] = (
                            round(meas_val / limit * 100, 2) if limit > 0 else 0.0)
        except Exception as e:
            print(f"[Warning] Oscilloscope get_measure_all query: {e}")

        if not meas_dict:
            for ch_num, ch_data in sorted(self.scope_channels.items()):
                if not ch_data.get('enabled', False):
                    continue
                name = self._channel_name(ch_num)
                try:
                    labels, values = self.oscilloscope.get_measure(ch_num)
                    meas_val = round(float(values[0]), 2) if values and values[0] is not None else 0.0
                except Exception:
                    meas_val = 0.0
                meas_dict[f"{name} Max"] = meas_val
                if ch_data.get('derating') is not None:
                    limit = ch_data.get('derating')
                    meas_dict[f"{name} Derating (%)"] = round(meas_val / limit * 100, 2) if limit > 0 else 0.0

        for ch_num, ch_data in sorted(self.scope_channels.items()):
            if not (ch_data.get('enabled', False) and (ch_data.get('cursor') or {}).get('enabled')):
                continue
            name = self._channel_name(ch_num)
            dx = dy = 0.0
            try:
                cursor_vals = self.oscilloscope.get_cursor(cursor=ch_num) or {}
                dx = round(float(cursor_vals.get('delta x', 0.0)), 6)
                dy = round(float(cursor_vals.get('delta y', 0.0)), 3)
            except Exception as e:
                print(f"[Warning] Cursor query error on CH{ch_num}: {e}")
            meas_dict[f"{name} dX"] = dx
            meas_dict[f"{name} Fsw (kHz)"] = round((1.0 / abs(dx)) / 1000.0, 2) if dx else 0.0
            meas_dict[f"{name} dY"] = dy
        return meas_dict

    def discharge_output(self, times: int = 2):
        """Discharge the outputs by pulsing the electronic loads with the input off."""
        if times <= 0:
            return
        i_discharge = float(getattr(self, 'DISCHARGE_CC_A', 1.0))
        try:
            self.input_supply.turn_off()
        except Exception:
            pass
        loads = [ld for ld in (self.electronic_load, self.electronic_load_2) if ld is not None]
        for _ in range(times):
            for ld in loads:
                try:
                    ld.set_cc_high(i_discharge)
                    ld.turn_on()
                except Exception:
                    pass
            sleep(1)
            for ld in loads:
                try:
                    ld.turn_off()
                except Exception:
                    pass
            sleep(0.5)
        for ld in loads:
            try:
                ld.reset_values()
            except Exception:
                pass

    def process_data_row(self):
        row = super().process_data_row()
        vin_coupling_col = 2
        if isinstance(row[vin_coupling_col], float) and row[vin_coupling_col] == int(row[vin_coupling_col]):
            row[vin_coupling_col] = int(row[vin_coupling_col])
        if isinstance(row[3], float) and row[3] == int(row[3]):
            row[3] = int(row[3])
        for i in range(4, len(row)):
            if isinstance(row[i], (int, float)):
                row[i] = round(float(row[i]), 2)
        return row

    @staticmethod
    def _excel_num_format_for_header(header: str) -> str:
        if header in ('Ch1 Load', 'Ch2 Load', 'Trig Ch', 'Waveform File'):
            return ''
        if header.endswith('(rms)') or header == 'Freq (Hz)':
            return '0'
        return '#,##0.00'

    def _format_main_sheet_data_rows(self, start_row: int, end_row: int):
        wb = openpyxl.load_workbook(self.data_file_path)
        ws = wb[self.sheet_name]
        for col_idx, header in enumerate(self.header_list, start=1):
            fmt = self._excel_num_format_for_header(header)
            if not fmt:
                continue
            for row in range(start_row, end_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
        scope_start = self.scope_excel_col()
        for offset, header in enumerate(self.scope_header_list):
            fmt = self._excel_num_format_for_header(header)
            if not fmt:
                continue
            col = scope_start + offset
            for row in range(start_row, end_row + 1):
                cell = ws.cell(row=row, column=col)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
        wb.save(self.data_file_path)
        wb.close()

    def _apply_main_sheet_number_formats(self):
        """Format any data rows already written (safety pass after multi-Vin export)."""
        try:
            wb = openpyxl.load_workbook(self.data_file_path)
            ws = wb[self.sheet_name]
            end_row = ws.max_row
            start_row = self.excel_first_data_row + 1
            if end_row >= start_row:
                wb.close()
                self._format_main_sheet_data_rows(start_row, end_row)
            else:
                wb.close()
        except Exception as e:
            print(f"[Warning] Excel number format pass: {e}")

    def _link_capture_main_sheet_rows(self):
        row = self.excel_first_data_row + 1
        for rec in self.captured_records:
            rec['main_sheet_row'] = row
            row += 1

    def _scaled_openpyxl_image(self, img_path: str):
        img = OpenpyxlImage(img_path)
        width_px, height_px = img.width, img.height
        try:
            from PIL import Image as PILImage
            with PILImage.open(img_path) as pim:
                width_px, height_px = pim.size
        except Exception:
            pass
        max_w = self.EMBED_IMAGE_MAX_WIDTH_PX
        scale = min(1.0, max_w / width_px) if width_px > max_w else 1.0
        img.width = int(width_px * scale)
        img.height = int(height_px * scale)
        return img

    def embed_waveforms_in_excel(self):
        if OpenpyxlImage is None or not self.captured_records:
            return
        try:
            wb = openpyxl.load_workbook(self.data_file_path)
            sheet = f"{self.short_title}_Waveforms"
            if sheet in wb.sheetnames:
                del wb[sheet]
            ws = wb.create_sheet(title=sheet)
            ws['B2'] = f"{self.title} - Captured Waveforms"
            ws['B2'].font = CellFont(size=14, bold=True, color="1F497D")
            ws.column_dimensions['B'].width = 28
            ws.column_dimensions['C'].width = 22

            current_row = 4
            label_font = CellFont(bold=True, size=11, color="2952A3")
            hdr_font = CellFont(bold=True, size=10, color="1F497D")
            for rec in self.captured_records:
                if not os.path.exists(rec['path']):
                    continue
                vin_disp = int(rec['vin']) if rec['vin'] == int(rec['vin']) else rec['vin']
                ws[f"B{current_row}"] = (
                    f"{vin_disp:g} V{self.coupling} | Ch1 {rec['ch1_lbl']} / Ch2 {rec['ch2_lbl']}")
                ws[f"B{current_row}"].font = label_font
                current_row += 1

                img = self._scaled_openpyxl_image(rec['path'])
                ws.add_image(img, f"B{current_row}")
                rows_for_img = max(12, int(img.height / 18) + 2)
                current_row += rows_for_img

                ws.cell(row=current_row, column=2, value="Parameter").font = hdr_font
                ws.cell(row=current_row, column=3, value="Value").font = hdr_font
                current_row += 1
                detail_rows = [
                    ("Main sheet row", rec.get('main_sheet_row')),
                    ("Efficiency (%)", round(rec['eff'], 2)),
                    (f"Trig Ch", f"CH{self.trigger_channel}"),
                ]
                for key, val in rec.get('measurements', {}).items():
                    if val not in (None, "N/A"):
                        detail_rows.append((key, val))
                detail_rows.append(("Waveform file", os.path.basename(rec['path'])))
                for param, val in detail_rows:
                    ws.cell(row=current_row, column=2, value=param)
                    cval = ws.cell(row=current_row, column=3, value=val)
                    if isinstance(val, float):
                        cval.number_format = '#,##0.00'
                    current_row += 1
                current_row += 1

            wb.save(self.data_file_path)
            wb.close()
        except Exception as e:
            print(f"Error embedding waveforms into Excel: {e}")
            print(traceback.format_exc())

    ###########################################################################
    #                      Results page / time estimate                       #
    ###########################################################################
    def define_output_data_objects(self):
        super().define_output_data_objects()
        self.primary_vds_name = None
        self.derating_limit_v = None
        for ch, ch_data in sorted(self.scope_channels.items()):
            if ch_data.get('enabled') and (ch_data.get('name') == 'Primary Vds' or self.primary_vds_name is None):
                self.primary_vds_name = ch_data.get('name') or f"CH{ch}"
                self.derating_limit_v = ch_data.get('derating')
                if ch_data.get('name') == 'Primary Vds':
                    break

        max_vin = max([v[0] for v in self.vin_list]) if self.vin_list else 265
        self.vds_vs_vin_plot = PlottableObject(
            title="Max Vds vs Input Voltage",
            type=PlotType.LINE,
            x_label=f"Input Voltage (V{self.coupling})",
            y_label="Maximum Vds (V)",
            x_range=(0, max_vin * 1.1),
            y_range=(0, 800),
            plot_series_list=[])
        for _, _, ch1_lbl, ch2_lbl in self.load_pairs:
            self.vds_vs_vin_plot.add_plot_series(name=f"Ch1 {ch1_lbl} / Ch2 {ch2_lbl}", x_values=[], y_values=[])
        if self.derating_limit_v:
            self.vds_vs_vin_plot.add_plot_series(
                name=f"Derating Limit ({self.derating_limit_v} V)", x_values=[], y_values=[])

        self.test_data_table = DataTable(header=self.header_list + self.scope_header_list, data=[])

    def update_output_data(self):
        td1 = self.test_data
        td2 = self.test_data_2
        self.load_vs_efficiency_plot.append_plot_data(plot_index=self.vin_index, x=td1.iout_A, y=self.eff_total_pct)
        self.ch1_vreg_plot.append_plot_data(plot_index=self.vin_index, x=td1.iout_A, y=td1.vout_V)
        self.ch2_vreg_plot.append_plot_data(plot_index=self.vin_index, x=td2.iout_A, y=td2.vout_V)

        scope_row = getattr(self, 'current_scope_row', [])
        if self.primary_vds_name:
            col = f"{self.primary_vds_name} Max"
            if col in self.scope_header_list:
                try:
                    vds = float(scope_row[self.scope_header_list.index(col)])
                    self.vds_vs_vin_plot.append_plot_data(
                        plot_index=self.current_load_index, x=self.vin_freq[0], y=vds)
                except (ValueError, TypeError, IndexError):
                    pass
        if self.derating_limit_v and self.current_load_index == 0:
            self.vds_vs_vin_plot.append_plot_data(
                plot_index=len(self.load_pairs), x=self.vin_freq[0], y=self.derating_limit_v)

        self.test_data_table.add_data_row(list(self.process_data_row()) + list(scope_row))
        self.test_data_update.emit([
            [self.load_vs_efficiency_plot, self.ch1_vreg_plot, self.ch2_vreg_plot, self.vds_vs_vin_plot],
            self.test_data_table])

    def estimate_remaining(self, vin_index_t, load_index_t, vin_delays: bool = True):
        soak = self.soak_time
        remaining_time_s = 0
        remaining_steps = 0
        start_adding_time = False
        vin_index = 0
        load_index = 0

        def add_time(t, add_step=False):
            nonlocal remaining_time_s, remaining_steps, start_adding_time
            if not start_adding_time and vin_index == vin_index_t and load_index == load_index_t:
                start_adding_time = True
            if start_adding_time:
                remaining_time_s += t
                if add_step:
                    remaining_steps += 1

        for vin_index, _ in enumerate(self.vin_list):
            load_index = 0
            if vin_delays:
                add_time(4)
                if self.usbpd_test:
                    add_time(1)
                add_time(soak.initial_soak if vin_index == 0 else soak.soak_per_line)
            for load_index, _ in enumerate(self.load_pairs):
                add_time(2, True)
                add_time(soak.soak_per_load)
                add_time(self.settle_time)
                add_time(soak.integration_time)
                add_time(1.7)
                add_time(12 if self.auto_find_trigger else 2)
                add_time(self.post_capture_delay)
            add_time(3)
        add_time(self.discharge_pulses * 1.5)
        return remaining_time_s, remaining_steps

    def _waveform_queue_time_line(self) -> str:
        import datetime
        total_s = getattr(self, 'total_time', 0) or 0
        total_txt = str(datetime.timedelta(seconds=round(total_s, 0)))
        if self.status in [TestStatus.STOPPED, TestStatus.FAILED, TestStatus.COMPLETE, TestStatus.SKIPPED]:
            return str(self.status)
        rem_txt = getattr(self, 'estimated_time_txt', '') or str(
            datetime.timedelta(seconds=round(self.estimated_time_s, 0)))
        prog = getattr(self, 'progress_txt', '') or '0%'
        return f"{self.status}: {rem_txt} remaining / {total_txt} total, {prog}"

    def update_test_list_text(self):
        super().update_test_list_text()
        lines = [
            ln for ln in self.test_list_text.split("\n")
            if ln and not ln.startswith("Oscilloscope Measurement:")
        ]
        extra = [f"Trigger: CH{self.trigger_channel} (Delta: {self.trigger_delta:g}V) | Prompts: {self.user_prompts}"]
        merged = lines[:2] + extra + lines[2:]
        time_line = self._waveform_queue_time_line()
        merged = [
            time_line if ln.startswith(f"{TestStatus.IN_QUEUE}:") or ln.startswith(f"{TestStatus.IN_PROGRESS}:")
            else ln
            for ln in merged
        ]
        self.test_list_text = "\n".join(merged)

    ###########################################################################
    #                         Save / load test items                          #
    ###########################################################################
    def get_dict(self) -> dict:
        d = super().get_dict()
        d['UNIT_ID'] = self.unit_id
        d['AMBIENT_TEMP'] = self.ambient_temp
        d['SCOPE_CHANNELS'] = self.scope_channels
        return d

    @staticmethod
    def extract_test_condition(test_item_dict: dict) -> TestConditions:
        tc = EfficiencyTest_2Port.extract_test_condition(test_item_dict)
        tc.unit_id = test_item_dict.get('UNIT_ID', '00')
        tc.ambient_temp = test_item_dict.get('AMBIENT_TEMP', 25.0)
        saved = test_item_dict.get('SCOPE_CHANNELS') or {}
        default = EfficiencyWaveformTest_2Port.tc_default.scope_channels
        tc.scope_channels = {
            ch: saved.get(ch, saved.get(str(ch), default[ch])) for ch in (1, 2, 3, 4)
        }
        return tc
