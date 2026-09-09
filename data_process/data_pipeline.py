import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from plotter.plotter import DataTable, PlottableObject, PlotSeries, PlotType

class TestTelemetryLogger:
    """
    Centralized data logger for ATE tests.
    Handles appending data to Excel worksheets, GUI DataTables, and PyQtGraph PlottableObjects.
    """
    def __init__(self, headers: List[str]):
        self.headers = headers
        self.data_table = DataTable(headers=headers)
        self.plottables: List[PlottableObject] = []
        
        # Excel references
        self.workbook: Optional[openpyxl.Workbook] = None
        self.worksheet: Optional[Worksheet] = None
        self.excel_file_path: Optional[str] = None
        self.current_excel_row = 5 # Starts writing data from row 5
        
    def setup_excel(self, file_path: str, sheet_name: str = "Raw Data"):
        """Initialize the Excel workbook and write headers."""
        self.excel_file_path = file_path
        self.workbook = openpyxl.Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = sheet_name
        
        # Write headers
        for col_idx, header in enumerate(self.headers, start=1):
            self.worksheet.cell(row=self.current_excel_row - 1, column=col_idx, value=header)
            
        self.workbook.save(self.excel_file_path)
        
    def add_plot(self, title: str, x_header: str, y_headers: List[str], plot_type: PlotType = PlotType.LINE_PLOT):
        """Register a new plot to be updated when data is recorded."""
        if x_header not in self.headers:
            raise ValueError(f"X header '{x_header}' not in defined headers.")
            
        series_list = []
        for y_header in y_headers:
            if y_header not in self.headers:
                raise ValueError(f"Y header '{y_header}' not in defined headers.")
            series_list.append(PlotSeries(x_data_name=x_header, y_data_name=y_header))
            
        plot = PlottableObject(title=title, plot_series=series_list, plot_type=plot_type)
        self.plottables.append(plot)

    def record_datapoint(self, data: Dict[str, Any]):
        """
        Record a single row of telemetry data.
        Updates the DataTable, Plottables, and Excel sheet simultaneously.
        """
        # Ensure data dict matches headers
        row_data = []
        for header in self.headers:
            val = data.get(header, "")
            row_data.append(val)
            
        # 1. Update GUI DataTable
        self.data_table.add_row(row_data)
        
        # 2. Update Plottables
        for plot in self.plottables:
            for series in plot.plot_series:
                x_val = data.get(series.x_data_name)
                y_val = data.get(series.y_data_name)
                if x_val is not None and y_val is not None and x_val != "" and y_val != "":
                    series.x_data.append(float(x_val))
                    series.y_data.append(float(y_val))
                    
        # 3. Update Excel Sheet
        if self.worksheet and self.workbook:
            for col_idx, val in enumerate(row_data, start=1):
                self.worksheet.cell(row=self.current_excel_row, column=col_idx, value=val)
            self.current_excel_row += 1
            
    def get_gui_update_payload(self) -> List[Any]:
        """Returns the payload expected by the test_data_update signal."""
        return [self.plottables, self.data_table]

    def close(self):
        """Finalize and close files."""
        if self.workbook:
            self.workbook.save(self.excel_file_path)
            self.workbook.close()
